from django.urls import path
from django.contrib import admin
from django.http import HttpResponse
from django.contrib.admin import SimpleListFilter
from django.db.models import Case, When, Value, IntegerField
from django.utils.translation import gettext_lazy as _
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

from .models import HistoryActionJournal


# =======================
# ORGANIZATION ORDER FILTER
# =======================
class OrganizationOrderFilter(SimpleListFilter):
    title = "Tashkilot faolligi bo'yicha"
    parameter_name = "org_order"

    def lookups(self, request, model_admin):
        return (
            ('active_first', 'Faol tashkilotlar birinchi'),
            ('inactive_first', 'Nofaol tashkilotlar birinchi'),
        )

    def queryset(self, request, queryset):
        if self.value() in ('active_first', 'inactive_first'):
            qs = queryset.annotate(
                org_active=Case(
                    When(
                        user__organization__user__last_login__isnull=False,
                        then=Value(1)
                    ),
                    default=Value(0),
                    output_field=IntegerField()
                )
            )
            if self.value() == 'active_first':
                return qs.order_by('-org_active', '-created_at')
            return qs.order_by('org_active', '-created_at')

        return queryset


# =======================
# ADMIN CONFIG
# =======================
@admin.register(HistoryActionJournal)
class HistoryActionJournalAdmin(admin.ModelAdmin):
    # -------- LIST VIEW ----------
    list_display = (
        'user',
        'user_organization',
        'action',
        'model_name',
        'object_repr',
        'change_message',
        'created_at',
    )

    # -------- FILTERS ----------
    list_filter = (
        'user',
        'user__organization',
        OrganizationOrderFilter,
        'model_name',
        'action',
        'created_at',
    )

    # -------- SEARCH ----------
    search_fields = (
        'user__email',
        'model_name',
        'object_repr',
        'change_message',
    )

    actions = ['export_to_excel']

    # -------- CUSTOM DISPLAY ----------
    def user_organization(self, obj):
        if obj.user and getattr(obj.user, 'organization', None):
            org = obj.user.organization
            return org.name_uz if hasattr(org, 'name_uz') and org.name_uz else str(org)
        return "-"

    user_organization.short_description = "Tashkilot"

    # -------- EXCEL GENERATOR ----------
    def _create_excel_response(self, rows, filename="tarix.xlsx"):
        """Excel fayl yaratish uchun umumiy metod"""
        wb = Workbook()
        ws = wb.active
        ws.title = "Tarix"

        # Stil uchun ranglar
        header_fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
        header_font = Font(bold=True, color="FFFFFF", size=11)
        header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

        # Border
        thin_border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )

        # Data alignment - wrap_text=False qilib qo'yamiz
        data_alignment_left = Alignment(horizontal="left", vertical="top", wrap_text=False)
        data_alignment_center = Alignment(horizontal="center", vertical="top", wrap_text=False)

        # Font for data
        data_font = Font(size=10)

        # Headerlar
        headers = [
            '№',
            'Foydalanuvchi',
            'Tashkilot',
            'Amal',
            'Model',
            'Obyekt',
            "O'zgarish xabari",
            'Sana'
        ]

        # Header qo'shish
        for col_num, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col_num)
            cell.value = header
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = header_alignment
            cell.border = thin_border

        # Ma'lumotlar qo'shish
        for row_num, obj in enumerate(rows, start=2):
            # №
            cell = ws.cell(row=row_num, column=1)
            cell.value = row_num - 1
            cell.alignment = data_alignment_center
            cell.border = thin_border
            cell.font = data_font

            # Foydalanuvchi
            cell = ws.cell(row=row_num, column=2)
            if obj.user:
                user_display = f"{obj.user.first_name} {obj.user.last_name}".strip()
                if not user_display:
                    user_display = obj.user.email if obj.user.email else "-"
                cell.value = user_display
            else:
                cell.value = "-"
            cell.alignment = data_alignment_left
            cell.border = thin_border
            cell.font = data_font

            # Tashkilot
            cell = ws.cell(row=row_num, column=3)
            if obj.user and getattr(obj.user, 'organization', None):
                org = obj.user.organization
                org_name = org.name_uz if hasattr(org, 'name_uz') and org.name_uz else str(org)
                cell.value = str(org_name) if org_name else "-"
            else:
                cell.value = "-"
            cell.alignment = data_alignment_left
            cell.border = thin_border
            cell.font = data_font

            # Amal
            cell = ws.cell(row=row_num, column=4)
            action_display = str(obj.get_action_display()) if hasattr(obj, 'get_action_display') else str(obj.action)
            cell.value = action_display
            cell.alignment = data_alignment_center
            cell.border = thin_border
            cell.font = data_font

            # Model
            cell = ws.cell(row=row_num, column=5)
            cell.value = str(obj.model_name) if obj.model_name else ''
            cell.alignment = data_alignment_left
            cell.border = thin_border
            cell.font = data_font

            # Obyekt
            cell = ws.cell(row=row_num, column=6)
            obyekt_text = str(obj.object_repr) if obj.object_repr else ''
            # Juda uzun matnlarni qisqartirish (ixtiyoriy)
            if len(obyekt_text) > 100:
                obyekt_text = obyekt_text[:97] + '...'
            cell.value = obyekt_text
            cell.alignment = data_alignment_left
            cell.border = thin_border
            cell.font = data_font

            # O'zgarish xabari
            cell = ws.cell(row=row_num, column=7)
            change_text = str(obj.change_message) if obj.change_message else ''
            # Juda uzun matnlarni qisqartirish (ixtiyoriy)
            if len(change_text) > 150:
                change_text = change_text[:147] + '...'
            cell.value = change_text
            cell.alignment = data_alignment_left
            cell.border = thin_border
            cell.font = data_font

            # Sana
            cell = ws.cell(row=row_num, column=8)
            cell.value = obj.created_at.strftime("%Y-%m-%d %H:%M:%S") if obj.created_at else ''
            cell.alignment = data_alignment_center
            cell.border = thin_border
            cell.font = data_font

            # Qator balandligini belgilash - har bir qator uchun
            ws.row_dimensions[row_num].height = 30

        # Ustunlar kengligini sozlash
        ws.column_dimensions['A'].width = 6  # №
        ws.column_dimensions['B'].width = 25  # Foydalanuvchi
        ws.column_dimensions['C'].width = 60  # Tashkilot - KENGAYTIRILDI
        ws.column_dimensions['D'].width = 12  # Amal
        ws.column_dimensions['E'].width = 18  # Model
        ws.column_dimensions['F'].width = 70  # Obyekt - KENGAYTIRILDI
        ws.column_dimensions['G'].width = 80  # O'zgarish xabari - KENGAYTIRILDI
        ws.column_dimensions['H'].width = 20  # Sana

        # Birinchi qatorni (header) balandligini sozlash
        ws.row_dimensions[1].height = 35

        # Freeze panes - header qatorini doimo ko'rinishda qoldirish
        ws.freeze_panes = 'A2'

        response = HttpResponse(
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
        response['Content-Disposition'] = f'attachment; filename="{filename}"'
        wb.save(response)
        return response

    # -------- ACTION: EXPORT SELECTED ----------
    def export_to_excel(self, request, queryset):
        """Tanlangan yozuvlarni Excel formatda yuklab olish"""
        return self._create_excel_response(queryset, filename="tanlangan_tarix.xlsx")

    export_to_excel.short_description = _("Excel formatda yuklab olish")

    # -------- CUSTOM URL: EXPORT ALL ----------
    def get_urls(self):
        urls = super().get_urls()
        custom_urls = [
            path(
                'export-all/',
                self.admin_site.admin_view(self.export_all_logs),
                name='history_export_all',
            ),
        ]
        return custom_urls + urls

    def export_all_logs(self, request):
        """Barcha yozuvlarni Excel formatda yuklab olish"""
        rows = HistoryActionJournal.objects.all()
        return self._create_excel_response(rows, filename="barcha_tarix.xlsx")

    # -------- ADMIN PAGE CONTEXT ----------
    def changelist_view(self, request, extra_context=None):
        extra_context = extra_context or {}
        extra_context['export_all_url'] = 'export-all/'
        return super().changelist_view(request, extra_context=extra_context)

    # -------- PERMISSIONS ----------
    def has_add_permission(self, request):
        return False

    def has_change_permission(self, request, obj=None):
        return False

    def has_delete_permission(self, request, obj=None):
        return False

    def has_view_permission(self, request, obj=None):
        return True