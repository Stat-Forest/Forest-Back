from apps.reports.exports.excel_export import purge_old_exports
from apps.utils.excel_safe import set_safe
import os
from datetime import datetime
from typing import Dict, Any, Tuple, List, Optional
from collections import defaultdict

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
from django.conf import settings
from django.utils.encoding import force_str


# ====================================================
# SAFE VALUE CONVERTER
# ====================================================
def _safe_excel_value(value):
    if value is None:
        return ""

    if isinstance(value, (int, float, bool)):
        return value

    try:
        return force_str(value)
    except Exception:
        return str(value)


# ====================================================
# TREE STRUCTURE FOR HEADER
# ====================================================
class HeaderNode:
    """Header tree uchun node"""

    def __init__(self, name):
        self.name = name
        self.children = []
        self.colspan = 1
        self.level = 0

    def add_child(self, child):
        self.children.append(child)

    def calculate_span(self):
        """Har bir node ning colspan ini hisoblash"""
        if not self.children:
            self.colspan = 1
        else:
            total = 0
            for child in self.children:
                child.calculate_span()
                total += child.colspan
            self.colspan = total
        return self.colspan

    def get_max_depth(self):
        """Maksimal chuqurlikni hisoblash"""
        if not self.children:
            return 1
        return 1 + max((child.get_max_depth() for child in self.children), default=0)


# ====================================================
# BUILD HEADER TREE FROM RESULT
# ====================================================
def _build_header_tree(result: Dict[str, Any], parent_node: Optional[HeaderNode] = None, path=None):
    """Result dict dan header tree qurish"""
    if path is None:
        path = []

    nodes = []

    for key, value in result.items():
        current_path = path + [key]

        if isinstance(value, dict):
            # Agar dict bo'lsa, rekursiv davom ettiramiz
            node = HeaderNode(key)
            child_nodes = _build_header_tree(value, node, current_path)
            for child in child_nodes:
                node.add_child(child)
            nodes.append(node)
        else:
            # Leaf node
            node = HeaderNode(key)
            nodes.append(node)

    return nodes


# ====================================================
# FLATTEN RESULT FOR VALUES
# ====================================================
def _flatten_result(result: Dict[str, Any], path=None):
    """Result dict dan qiymatlarni yig'ish (leaf node lar uchun)"""
    if path is None:
        path = []

    rows = {}

    for key, value in result.items():
        if isinstance(value, dict):
            # Rekursiv davom ettiramiz
            sub_rows = _flatten_result(value, path + [key])
            rows.update(sub_rows)
        else:
            # Leaf node - path + key dan unique identifier yaratamiz
            full_path = path + [key]
            # Key sifatida tuple ishlatamiz (path)
            path_key = tuple(full_path)
            rows[path_key] = _safe_excel_value(value)

    return rows


# ====================================================
# COLLECT ALL PATHS FROM NODES
# ====================================================
def _collect_paths(nodes: List[HeaderNode], current_path=None):
    """Barcha leaf node larga boradigan path larni yig'ish"""
    if current_path is None:
        current_path = []

    paths = []

    for node in nodes:
        new_path = current_path + [node.name]
        if not node.children:
            paths.append(new_path)
        else:
            paths.extend(_collect_paths(node.children, new_path))

    return paths


# ====================================================
# WRITE HEADER ROWS
# ====================================================
def _write_header(ws, root_nodes: List[HeaderNode], start_row: int, start_col: int):
    """Header tree ni Excel ga yozish"""

    # Har bir node ning colspan ini hisoblaymiz
    for node in root_nodes:
        node.calculate_span()

    # Maksimal chuqurlikni aniqlaymiz
    max_depth = max((node.get_max_depth() for node in root_nodes), default=1)

    # Header grid yaratish
    header_grid = [["" for _ in range(100)] for _ in range(max_depth)]

    # Recursive function to place nodes
    def place_node(node, level, col):
        header_grid[level][col] = node.name
        if node.children:
            # Merge cells for this node across its children
            ws.merge_cells(
                start_row=start_row + level,
                start_column=start_col + col + 1,
                end_row=start_row + level,
                end_column=start_col + col + node.colspan
            )
            current_col = col
            for child in node.children:
                current_col = place_node(child, level + 1, current_col)
            return col + node.colspan
        else:
            # Leaf node - merge vertically if needed
            if level < max_depth - 1:
                ws.merge_cells(
                    start_row=start_row + level,
                    start_column=start_col + col + 1,
                    end_row=start_row + max_depth - 1,
                    end_column=start_col + col + 1
                )
            return col + 1

    # Place all root nodes
    col_index = 0
    for node in root_nodes:
        col_index = place_node(node, 0, col_index)

    # Write header grid to worksheet
    for level in range(max_depth):
        for col in range(col_index):
            value = header_grid[level][col]
            if value:  # Only write non-empty cells
                cell = ws.cell(row=start_row + level, column=start_col + col + 1)
                set_safe(cell, value)

    return max_depth, col_index


# ====================================================
# MAIN EXPORT FUNCTION
# ====================================================
def export_monitoring_to_excel(data: Dict[str, Any], lang: str = "uz") -> Tuple[str, str]:
    now = datetime.now()
    filename = f"forest_monitoring_{now.strftime('%Y%m%d_%H%M%S')}.xlsx"

    export_dir = getattr(settings, "EXPORT_ROOT", os.path.join(settings.BASE_DIR, "exports"))
    os.makedirs(export_dir, exist_ok=True)
    purge_old_exports(export_dir)
    file_path = os.path.join(export_dir, filename)

    wb = Workbook()
    ws = wb.active
    ws.title = "Monitoring"

    # ====================================================
    # EXPORT TIME
    # ====================================================
    ws.append([f"Export qilingan vaqt: {now.strftime('%Y-%m-%d %H:%M:%S')}"])
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=5)
    ws["A1"].font = Font(bold=True)

    current_row = 3  # Header 3-qatordan boshlanadi

    # ====================================================
    # BUILD HEADER TREE FROM REPUBLIC RESULT
    # ====================================================
    result_data = data.get("result", {})
    header_nodes = _build_header_tree(result_data)

    # Fixed columns: Level va Hudud
    fixed_cols = 2

    # Write header
    header_depth, header_width = _write_header(ws, header_nodes, current_row, fixed_cols)

    # Write fixed headers
    ws.cell(row=current_row, column=1, value="Level")
    ws.cell(row=current_row, column=2, value="Hudud")

    # Merge fixed headers vertically if needed
    if header_depth > 1:
        ws.merge_cells(
            start_row=current_row,
            start_column=1,
            end_row=current_row + header_depth - 1,
            end_column=1
        )
        ws.merge_cells(
            start_row=current_row,
            start_column=2,
            end_row=current_row + header_depth - 1,
            end_column=2
        )

    # Style header
    header_fill = PatternFill("solid", fgColor="D9E1F2")
    thin_border = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin"),
    )

    for row in range(current_row, current_row + header_depth):
        for col in range(1, fixed_cols + header_width + 1):
            cell = ws.cell(row=row, column=col)
            cell.font = Font(bold=True)
            cell.fill = header_fill
            cell.border = thin_border
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    # ====================================================
    # COLLECT ALL PATHS FOR DATA ROWS
    # ====================================================
    all_paths = _collect_paths(header_nodes)

    # ====================================================
    # WRITE DATA ROWS
    # ====================================================
    def write_node(node, row_num):
        level = _safe_excel_value(node.get("level"))
        hudud = _safe_excel_value(node.get("hudud"))

        # Qiymatlarni flatten qilish
        flattened = _flatten_result(node.get("result", {}))

        # Row yozish
        ws.cell(row=row_num, column=1, value=level)
        set_safe(ws.cell(row=row_num, column=2), hudud)

        # Har bir path uchun qiymatni topish
        for col_idx, path in enumerate(all_paths):
            path_key = tuple(path)
            value = flattened.get(path_key, "")
            set_safe(ws.cell(row=row_num, column=fixed_cols + col_idx + 1), value)

        next_row = row_num + 1

        # Children lar uchun rekursiv
        for child in node.get("children", []):
            next_row = write_node(child, next_row)

        return next_row

    data_start_row = current_row + header_depth
    write_node(data, data_start_row)

    # ====================================================
    # STYLE DATA
    # ====================================================
    for row in range(data_start_row, ws.max_row + 1):
        for col in range(1, fixed_cols + header_width + 1):
            cell = ws.cell(row=row, column=col)
            cell.border = thin_border
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    # ====================================================
    # AUTO WIDTH
    # ====================================================
    for col in range(1, fixed_cols + header_width + 1):
        letter = get_column_letter(col)
        max_length = 0
        for row in range(1, ws.max_row + 1):
            val = ws.cell(row=row, column=col).value
            if val:
                max_length = max(max_length, len(str(val)))
        ws.column_dimensions[letter].width = min(max(10, max_length + 2), 40)

    wb.save(file_path)
    return file_path, filename