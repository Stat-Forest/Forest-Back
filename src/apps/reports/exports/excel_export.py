from apps.utils.excel_safe import append_safe
import os
import time
from datetime import datetime
from decimal import Decimal, InvalidOperation
from typing import List, Dict, Any, Tuple, Optional, Set

from django.conf import settings
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

from apps.reports.services.get_two_numbers_in_percent import (
    decimal_to_display_string,
    decimal_to_db_string,
    safe_decimal,
)


def apply_border_to_range(ws, min_row, max_row, min_col, max_col, border):
    for r in range(min_row, max_row + 1):
        for c in range(min_col, max_col + 1):
            ws.cell(row=r, column=c).border = border


def _ordering_key(ordering: str) -> tuple:
    return tuple(int(p) if p.isdigit() else p for p in (ordering or "").split("."))


def _is_numeric_like(value) -> bool:
    """
    Qiymat numeric ko'rinishdami yoki yo'qmi tekshiradi.

    Muhim:
    Bu yerda safe_decimal() ishlatmaymiz, chunki safe_decimal noto'g'ri textlarda
    ham Decimal("0") qaytarishi mumkin. Bu esa text qiymatlarni 0 qilib yuboradi.
    """
    if value is None:
        return False

    text = str(value).strip().replace(" ", "").replace(",", ".")
    if not text:
        return False

    try:
        Decimal(text)
        return True
    except (InvalidOperation, ValueError, TypeError):
        return False


def _format_excel_display_value(value) -> str:
    """
    Excelga yozishdan oldin numeric qiymatlarni vergul formatga o'tkazadi.

    DB/calculation:
        50.20

    Excel display:
        50,20

    Text qiymatlar o'z holicha qoladi.
    """
    if value is None:
        return ""

    text = str(value).strip()
    if not text:
        return ""

    if _is_numeric_like(text):
        return decimal_to_display_string(text)

    return text


def _non_empty_inners(inners) -> list:
    """
    Bo'sh qiymatli inner'larni olib tashlaydi.

    Nima uchun: oddiy (plantsiz) hisobotda qator raqami shu ro'yxatning
    indeksidan olinadi. Agar 2-ilova ustunlarida oldin bo'sh inner'lar
    turgan bo'lsa, qiymat pastki qatorga surilib ketardi — Excelda
    1-ilova tepada, 2-ilova pastda "zinapoya" bo'lib ko'rinardi.
    """
    return [
        inn for inn in (inners or [])
        if str((inn or {}).get("value") or "").strip() != ""
    ]


def _pick_best_value(inners: list) -> str:
    """
    Multiple inner qiymatlarni bitta Excel cell qiymatiga aylantiradi.

    Rules:
    - If no values -> empty string.
    - If one numeric value -> comma display format.
    - If multiple numeric-like values exist -> sum all numeric-like values and comma format.
    - Else -> return the first non-empty textual value.
    """
    if not inners:
        return ""

    values = [str((it.get("value") or "")).strip() for it in inners]

    if len(values) == 1:
        return _format_excel_display_value(values[0])

    numeric_values = []
    for value in values:
        if _is_numeric_like(value):
            numeric_values.append(safe_decimal(value))

    if numeric_values:
        total = sum(numeric_values, Decimal("0"))
        return _format_excel_display_value(total)

    return next((v for v in values if v), "")


def _sum_decimal_values(values) -> tuple[Decimal, bool]:
    """
    Iterable qiymatlarni Decimal sifatida yig'adi.

    returns:
        (total, seen)
    """
    total = Decimal("0")
    seen = False

    for value in values:
        if value in (None, ""):
            continue

        if not _is_numeric_like(value):
            continue

        total += safe_decimal(value)
        seen = True

    return total, seen


def _accumulate_decimal_string(current: Optional[str], new_val: Any) -> str:
    """
    Plant-mode uchun qiymatlarni yig'adi.
    Ichki hisob-kitob DB formatda, ya'ni nuqta bilan saqlanadi.
    Excelga yozish paytida vergulga aylantiriladi.
    """
    if new_val in (None, ""):
        return current or ""

    if current in (None, ""):
        return decimal_to_db_string(new_val) if _is_numeric_like(new_val) else str(new_val)

    if _is_numeric_like(current) and _is_numeric_like(new_val):
        total = safe_decimal(current) + safe_decimal(new_val)
        return decimal_to_db_string(total)

    return str(new_val) or str(current)


_INVALID_SHEET_CHARS = set(r"[]:*?/\'")


def _sheet_title(raw_title: str, used: Set[str], fallback: str) -> str:
    """
    Excel varaq nomi: 31 belgidan oshmaydi, taqiqlangan belgilarsiz,
    takrorlanmaydigan.
    """
    title = "".join(" " if ch in _INVALID_SHEET_CHARS else ch for ch in str(raw_title or ""))
    title = " ".join(title.split()).strip()[:31] or fallback

    candidate = title
    counter = 2
    while candidate.lower() in used:
        suffix = f" ({counter})"
        candidate = f"{title[:31 - len(suffix)]}{suffix}"
        counter += 1

    used.add(candidate.lower())
    return candidate


EXPORT_RETENTION_SECONDS = 24 * 3600


def purge_old_exports(export_dir: str, max_age: int = EXPORT_RETENTION_SECONDS) -> None:
    """Eksport fayllari diskda to'planib qolmasin (audit C6): 24 soatdan eski
    .xlsx fayllar yangi eksport paytida o'chiriladi (yuklab olish havolasi
    baribir shu muddatgacha ishlatiladi)."""
    try:
        now = time.time()
        for name in os.listdir(export_dir):
            if not name.lower().endswith(".xlsx"):
                continue
            path = os.path.join(export_dir, name)
            try:
                if os.path.isfile(path) and now - os.path.getmtime(path) > max_age:
                    os.remove(path)
            except OSError:
                pass
    except OSError:
        pass


def export_universal_report_to_excel_v2(
    reports: List[Dict[str, Any]],
    *,
    filename_prefix: str = "report_export",
    lang: str = "uz",
    statement_id: Optional[str] = None,
    aggregated_mode: bool = False,
) -> Tuple[str, str]:
    """
    Hisobotlarni Excelga chiqaradi.

    HAR BIR ILOVA — ALOHIDA VARAQ (sheet). Saytda ham har bir ilova alohida
    jadval bo'lib turadi, shuning uchun bitta varaqqa hammasini yonma-yon
    tiqishtirmaymiz: ilovalarda o'simliklar har xil bo'lgani sababli
    (1-ilovada ko'chat turlari, 2-ilovada nihol turlari) yagona varaqda
    katta bo'sh bloklar hosil bo'lardi.
    """
    from apps.reports.models import StatementColumn

    now = datetime.now()
    filename = f"{filename_prefix}_{now.strftime('%Y%m%d_%H%M%S')}.xlsx"
    export_dir = getattr(settings, "EXPORT_ROOT", os.path.join(settings.BASE_DIR, "exports"))
    os.makedirs(export_dir, exist_ok=True)
    purge_old_exports(export_dir)
    file_path = os.path.join(export_dir, filename)

    wb = Workbook()

    export_text = f"Export qilingan vaqt: {now.strftime('%Y-%m-%d %H:%M:%S')}"

    HEADER_START_ROW = 2

    thin_border = Border(
        left=Side(style="thin", color="000000"),
        right=Side(style="thin", color="000000"),
        top=Side(style="thin", color="000000"),
        bottom=Side(style="thin", color="000000"),
    )
    header_fill = PatternFill("solid", fgColor="D9E1F2")

    def _empty_file():
        ws = wb.active
        ws.title = "Reports"
        append_safe(ws, ["№", "Tashkilot nomi"])
        wb.save(file_path)
        return file_path, filename

    # =============================
    # Collect column IDs
    # =============================
    needed_ids: Set[str] = {
        str(col["statement_column_id"])
        for report in reports
        for item in report.get("report_items") or []
        for col in item.get("statement_columns") or []
        if col.get("statement_column_id")
    }

    if not needed_ids and not statement_id:
        return _empty_file()

    if not statement_id:
        statement_id = str(
            StatementColumn.objects.filter(id=next(iter(needed_ids)))
            .values_list("statement_id", flat=True)
            .first()
        )

    all_cols = list(
        StatementColumn.objects.filter(statement_id=statement_id)
        .select_related("parent")
        .only("id", "parent_id", "ordering", "title_uz", "title_ru", "is_organization")
    )
    id_map = {str(c.id): c for c in all_cols}

    def is_org(c):
        return getattr(c, "is_organization", False)

    # Jadval saytda qanday ko'rinsa, Excelda ham shunday bo'lishi kerak:
    # ustunlar ro'yxati hisobotdagi qiymatlardan emas, jadvalning O'ZIDAN
    # olinadi. Aks holda qiymati yo'q ustun (masalan hali hisoblanmagan
    # "Jami") Exceldan butunlay tushib qolardi.
    parent_ids = {str(c.parent_id) for c in all_cols if c.parent_id}
    export_ids = {
        str(c.id)
        for c in all_cols
        if not is_org(c) and str(c.id) not in parent_ids  # faqat leaf ustunlar
    }

    included_ids = set(export_ids)
    for cid in list(export_ids):
        cur = id_map.get(cid)
        while cur and cur.parent_id:
            included_ids.add(str(cur.parent_id))
            cur = id_map.get(str(cur.parent_id))

    children: Dict[Optional[str], List[Any]] = {}
    for cid in included_ids:
        c = id_map.get(cid)
        if not c or is_org(c):
            continue
        children.setdefault(str(c.parent_id) if c.parent_id else None, []).append(c)

    for k in children:
        children[k].sort(key=lambda x: _ordering_key(x.ordering or ""))

    def dfs(node):
        ch = children.get(str(node.id), [])
        if not ch:
            return [node] if str(node.id) in export_ids else []
        return [x for c in ch for x in dfs(c)]

    def depth(n):
        ch = children.get(str(n.id), [])
        return 1 if not ch else 1 + max(depth(c) for c in ch)

    def title_of(c):
        return (c.title_ru if lang == "ru" else c.title_uz) or c.title_uz

    roots = children.get(None, [])
    if not roots:
        return _empty_file()

    # "Tashkilot Jami" hisobotlari (individual_general) — har bir varaqda kerak
    ind_general_by_org: Dict[str, Dict[str, Any]] = {}
    for rep in reports:
        rep_items = rep.get("report_items") or []
        if not rep_items:
            continue
        if (rep_items[0] or {}).get("organization_id") == "individual":
            org_uuid = str(rep.get("organization")) if rep.get("organization") else None
            if org_uuid:
                ind_general_by_org[org_uuid] = rep

    # ==================================================================
    #  BITTA ILOVANI BITTA VARAQQA CHIZISH
    # ==================================================================
    def render_appendix(ws, root):
        leaf_ids = [str(c.id) for c in dfs(root)]
        leaf_id_set = set(leaf_ids)
        max_depth = depth(root)

        # Shu ilovada o'simlik kesimida qiymat bormi?
        any_plant = any(
            inn.get("plant_id")
            for report in reports
            for item in report.get("report_items") or []
            for col in item.get("statement_columns") or []
            if str(col.get("statement_column_id")) in leaf_id_set
            for inn in col.get("inners") or []
        )

        fixed_cols = 3 if any_plant else 2
        total_cols = fixed_cols + len(leaf_ids)

        # ---------- Export vaqti ----------
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=2)
        cell = ws.cell(row=1, column=1)
        cell.value = export_text
        cell.font = Font(bold=True, italic=True)
        cell.alignment = Alignment(horizontal="left", vertical="center")
        cell.border = thin_border

        # ---------- Sarlavha ----------
        header_grid = [["" for _ in range(total_cols)] for _ in range(max_depth)]
        header_grid[0][0] = "№"
        header_grid[0][1] = "Tashkilot nomi" if lang == "uz" else "Наименование организации"
        if any_plant:
            from apps.reports.models import Statement
            stmt = Statement.objects.filter(id=statement_id).first()
            stmt_name = (stmt.name_uz or stmt.name or "").lower() if stmt else ""
            if "daromad" in stmt_name:
                header_grid[0][2] = "Daromad manbalari" if lang == "uz" else "Источники дохода"
            else:
                header_grid[0][2] = "O'simlik nomi" if lang == "uz" else "Название растения"

        span_cache = {}

        def span(n):
            if n.id in span_cache:
                return span_cache[n.id]
            ch = children.get(str(n.id), [])
            span_cache[n.id] = 1 if not ch else sum(span(c) for c in ch)
            return span_cache[n.id]

        merges = []

        def place(node, level, col):
            header_grid[level][col] = title_of(node)
            sp = span(node)
            ch = children.get(str(node.id), [])

            if not ch:
                if level < max_depth - 1:
                    merges.append((
                        HEADER_START_ROW + level, col + 1,
                        HEADER_START_ROW + max_depth - 1, col + 1,
                    ))
                return col + 1

            merges.append((
                HEADER_START_ROW + level, col + 1,
                HEADER_START_ROW + level, col + sp,
            ))

            cur = col
            for c in ch:
                cur = place(c, level + 1, cur)
            return col + sp

        place(root, 0, fixed_cols)

        for row in header_grid:
            append_safe(ws, row)

        for fixed_col in range(1, fixed_cols + 1):
            ws.merge_cells(
                start_row=HEADER_START_ROW, start_column=fixed_col,
                end_row=HEADER_START_ROW + max_depth - 1, end_column=fixed_col,
            )

        for row in ws.iter_rows(
            min_row=HEADER_START_ROW, max_row=HEADER_START_ROW + max_depth - 1,
            min_col=1, max_col=total_cols,
        ):
            for cell in row:
                cell.font = Font(bold=True)
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                cell.fill = header_fill
                cell.border = thin_border

        for r1, c1, r2, c2 in merges:
            ws.merge_cells(start_row=r1, start_column=c1, end_row=r2, end_column=c2)

        apply_border_to_range(
            ws, HEADER_START_ROW, HEADER_START_ROW + max_depth - 1, 1, total_cols, thin_border
        )

        # DIQQAT: ws.cell(...) chaqirilsa, o'sha katak YARATILADI va ws.max_row
        # bittaga oshib ketadi — sarlavha bilan ma'lumot orasida bo'sh qator
        # paydo bo'lardi. Shuning uchun manzilni matn ko'rinishida beramiz.
        ws.freeze_panes = f"{get_column_letter(fixed_cols + 1)}{HEADER_START_ROW + max_depth}"

        data_start_row = ws.max_row + 1

        # ---------- Aggregated: Viloyat Jami / Respublika Jami ----------
        if aggregated_mode and not any_plant:
            for report in reports:
                rep_items = report.get("report_items") or []
                if not rep_items:
                    continue

                org_id_tag = (rep_items[0] or {}).get("organization_id")
                org_name = report.get("organization_name", "")
                row_no = report.get("row_number") or ""

                is_general_rep = bool(report.get("is_general") or org_id_tag == "general")
                is_republic_rep = bool(report.get("is_republic") or org_id_tag == "republic")

                if is_general_rep or is_republic_rep:
                    values = {}
                    for item in rep_items:
                        for col in item.get("statement_columns") or []:
                            cid = str(col.get("statement_column_id"))
                            if cid in leaf_id_set:
                                values[cid] = _pick_best_value(col.get("inners") or [])

                    row_values = [row_no, org_name] + [values.get(cid, "") for cid in leaf_ids]
                    if not any(v not in (None, "") for v in row_values[2:]):
                        continue

                    append_safe(ws, row_values)
                    for cell in ws[ws.max_row]:
                        cell.font = Font(bold=True)
                        cell.border = thin_border
                        cell.alignment = Alignment(
                            horizontal="center", vertical="center", wrap_text=True
                        )
                    continue

                per_col_inners: Dict[str, list] = {cid: [] for cid in leaf_ids}
                for item in rep_items:
                    for col in item.get("statement_columns") or []:
                        cid = str(col.get("statement_column_id"))
                        if cid in per_col_inners:
                            per_col_inners[cid] = _non_empty_inners(col.get("inners"))

                row_count = max((len(per_col_inners.get(cid, [])) for cid in leaf_ids), default=0)

                for r_idx in range(row_count):
                    row_values = [row_no, org_name]
                    for cid in leaf_ids:
                        inners = per_col_inners.get(cid, [])
                        row_values.append(
                            _format_excel_display_value((inners[r_idx] or {}).get("value"))
                            if r_idx < len(inners) else ""
                        )

                    if not any(v not in (None, "") for v in row_values[2:]):
                        continue
                    append_safe(ws, row_values)

            _finish_sheet(ws, data_start_row, total_cols, thin_border, bold_jami=True)
            return

        # ---------- Normal: hisobotlar + Tashkilot Jami ----------
        display_index = 0

        for report in reports:
            rep_items = report.get("report_items") or []
            if not rep_items:
                continue

            org_uuid = str(report.get("organization")) if report.get("organization") else None
            org_name = report.get("organization_name", "")
            if (rep_items[0] or {}).get("organization_id") == "individual":
                continue

            display_index += 1
            jami_label = "Tashkilot Jami" if lang == "uz" else "Итого по организации"

            # ----- Plantsiz (oddiy) ilova -----
            if not any_plant:
                per_col_inners: Dict[str, list] = {cid: [] for cid in leaf_ids}
                for item in rep_items:
                    for col in item.get("statement_columns") or []:
                        cid = str(col.get("statement_column_id"))
                        if cid in per_col_inners:
                            per_col_inners[cid] = _non_empty_inners(col.get("inners"))

                row_count = max((len(per_col_inners.get(cid, [])) for cid in leaf_ids), default=0)

                for r_idx in range(row_count):
                    row_values = [display_index, org_name]
                    for cid in leaf_ids:
                        inners = per_col_inners.get(cid, [])
                        row_values.append(
                            _format_excel_display_value((inners[r_idx] or {}).get("value"))
                            if r_idx < len(inners) else ""
                        )

                    # Butunlay bo'sh qator yozilmasin — saytda bunday qator yo'q
                    if not any(v not in (None, "") for v in row_values[2:]):
                        continue
                    append_safe(ws, row_values)

                total_row = [display_index, jami_label]
                totals_by_col = _totals_from_ind_report(
                    ind_general_by_org.get(org_uuid), leaf_id_set
                )

                for cid in leaf_ids:
                    value = totals_by_col.get(cid)
                    if value not in (None, ""):
                        total_row.append(_format_excel_display_value(value))
                        continue

                    col_total, seen = _sum_decimal_values(
                        [(inn or {}).get("value") for inn in per_col_inners.get(cid, [])]
                    )
                    total_row.append(_format_excel_display_value(col_total) if seen else "")

                if any(v not in (None, "") for v in total_row[2:]):
                    append_safe(ws, total_row)
                continue

            # ----- Plantlik ilova -----
            plant_names: Dict[str, str] = {}
            per_plant: Dict[str, Dict[str, str]] = {}
            # O'simliksiz qiymatlar: 3-ilova (svodka) shunday saqlanadi —
            # 1- va 2-ilovadan hisoblanib, tashkilot bo'yicha bitta qiymat
            # sifatida yoziladi (plant_id = NULL).
            plantless: Dict[str, str] = {}

            for item in rep_items:
                for col in item.get("statement_columns") or []:
                    cid = str(col.get("statement_column_id"))
                    if cid not in leaf_id_set:
                        continue

                    for inn in (col.get("inners") or []):
                        pid = inn.get("plant_id")
                        if not pid:
                            plantless[cid] = _accumulate_decimal_string(
                                plantless.get(cid), inn.get("value")
                            )
                            continue

                        plant_names[pid] = inn.get("plant_name") or plant_names.get(pid) or ""
                        prev = per_plant.setdefault(pid, {}).get(cid)
                        per_plant[pid][cid] = _accumulate_decimal_string(prev, inn.get("value"))

            for pid, pname in plant_names.items():
                vals = per_plant.get(pid, {})
                row = [display_index, org_name, pname]
                row += [_format_excel_display_value(vals.get(cid, "")) for cid in leaf_ids]

                if not any(v not in (None, "") for v in row[3:]):
                    continue
                append_safe(ws, row)

            # Tashkilot darajasidagi qiymat. Tartib:
            #   1) "Tashkilot Jami" hisobotidagi qiymat
            #   2) shu hisobotning o'simliksiz qiymati (svodka)
            #   3) o'simliklar bo'yicha yig'indi
            totals_by_col = _totals_from_ind_report(
                ind_general_by_org.get(org_uuid), leaf_id_set
            )

            jami_row = [display_index, jami_label, ""]
            for cid in leaf_ids:
                value = totals_by_col.get(cid)

                if value in (None, ""):
                    value = plantless.get(cid)

                if value in (None, ""):
                    total, seen = _sum_decimal_values(
                        [per_plant.get(pid, {}).get(cid) for pid in plant_names.keys()]
                    )
                    value = total if seen else ""

                jami_row.append(_format_excel_display_value(value))

            if any(v not in (None, "") for v in jami_row[3:]):
                append_safe(ws, jami_row)

        _finish_sheet(ws, data_start_row, total_cols, thin_border)

    # ==================================================================
    used_titles: Set[str] = set()
    for index, root in enumerate(roots):
        ws = wb.active if index == 0 else wb.create_sheet()
        ws.title = _sheet_title(title_of(root), used_titles, f"Ilova {index + 1}")
        render_appendix(ws, root)

    wb.save(file_path)
    return file_path, filename


def _totals_from_ind_report(ind_rep, leaf_id_set) -> Dict[str, str]:
    """"Tashkilot Jami" hisobotidan ustun bo'yicha qiymatlarni oladi."""
    totals: Dict[str, str] = {}
    if not ind_rep:
        return totals

    for item in ind_rep.get("report_items") or []:
        for col in item.get("statement_columns") or []:
            cid = str(col.get("statement_column_id"))
            if cid in leaf_id_set:
                inners = [
                    inn for inn in (col.get("inners") or [])
                    if not inn.get("plant_id")
                ]
                totals[cid] = _pick_best_value(inners)
    return totals


def _finish_sheet(ws, data_start_row, total_cols, thin_border, bold_jami=False):
    """Ma'lumot kataklariga chegara/tekislash va ustun kengliklari."""
    for row in ws.iter_rows(
        min_row=data_start_row, max_row=ws.max_row, min_col=1, max_col=total_cols
    ):
        for cell in row:
            cell.border = thin_border
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

            if bold_jami:
                org_name_cell_value = str(ws.cell(row=cell.row, column=2).value or "").lower()
                if "viloyat jami" in org_name_cell_value or "respublika jami" in org_name_cell_value:
                    cell.font = Font(bold=True)

    for col in range(1, ws.max_column + 1):
        letter = get_column_letter(col)

        if col == 1:
            ws.column_dimensions[letter].width = 5
            continue

        widest = max(
            (len(str(ws.cell(r, col).value or "")) for r in range(1, ws.max_row + 1)),
            default=0,
        )
        ws.column_dimensions[letter].width = min(max(8, widest + 1), 28)
