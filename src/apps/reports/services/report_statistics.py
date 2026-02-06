from django.db.models import Count
from openpyxl.styles import Font, Alignment

from apps.reports.models import Report


def get_report_statistics(filters=None):
    """
    Returns aggregated ERP report statistics:
    organization | statement | status | count
    """

    qs = Report.objects.all()

    if filters:
        if filters.get("organization"):
            qs = qs.filter(organization_id=filters["organization"])

        if filters.get("statement"):
            qs = qs.filter(statement_id=filters["statement"])

        if filters.get("status"):
            qs = qs.filter(status=filters["status"])

        if filters.get("deadline"):
            qs = qs.filter(deadline_id=filters["deadline"])

    return (
        qs.values(
            "organization__name",
            "statement__name_uz",
            "status",
        )
        .annotate(total=Count("id"))
        .order_by(
            "organization__name",
            "statement__name_uz",
            "status",
        )
    )


from openpyxl import Workbook
from openpyxl.utils import get_column_letter

def export_statistics_excel(filters=None):
    wb = Workbook()
    ws = wb.active
    ws.title = "Report Statistics"

    headers = ["Organization", "Statement", "Status", "Count"]
    ws.append(headers)

    # Header style
    for col in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col)
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center")

    # Data rows
    for row in get_report_statistics(filters):
        ws.append([
            row["organization__name"],
            row["statement__name_uz"],
            row["status"],
            row["total"],
        ])

    # Auto column width
    for col in ws.columns:
        max_length = 0
        col_letter = get_column_letter(col[0].column)

        for cell in col:
            if cell.value:
                max_length = max(max_length, len(str(cell.value)))

        ws.column_dimensions[col_letter].width = max_length + 6

    # Count column centered
    for cell in ws["D"]:
        cell.alignment = Alignment(horizontal="center")

    return wb