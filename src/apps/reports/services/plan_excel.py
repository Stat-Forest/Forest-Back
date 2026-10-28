# apps/reports/services/plan_excel.py
"""
Reja (Plan) uchun Excel shablon yaratish va yuklangan Excel faylni
PlanSerializer input formatiga o'girish.

Dizayn qarori:
    Har bir ustun sarlavhasi ostida (2-QATOR, yashirilgan) shu ustunning
    StatementColumn UUID'i saqlanadi. Import paytida sarlavha MATNIGA emas,
    aynan shu UUID'ga qarab column aniqlanadi. Shu sababli:
      - foydalanuvchi sarlavha matnini bir oz o'zgartirib qo'ysa ham import
        buzilmaydi;
      - uz/ru sarlavhalar orasidagi farq muammo tug'dirmaydi;
      - bir xil nomli (lekin turli statement/parent'ga tegishli) ustunlar
        aralashib ketmaydi.

Plant-plan statement'larda ham xuddi shunday: A-ustunda Plant ID saqlanadi
(import shu ID orqali qatorni plant bilan bog'laydi), lekin bu ustun
foydalanuvchiga ko'rinmasligi uchun YASHIRILGAN (hidden). Foydalanuvchi
faqat o'qiladigan "O'simlik nomi" ustunini ko'radi.

Foydalanuvchi 2-qatorni (ID qatori) yoki A-ustunni (Plant ID) o'chirib
yubormasligi kerak — shuning uchun ular `hidden` qilib qo'yiladi, lekin
fayldan butunlay olib tashlanmaydi.
"""

import openpyxl

from apps.utils.excel_safety import check_xlsx_zip
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from apps.reports.models import Plant, StatementColumn
from apps.reports.services.plan_columns import get_plan_input_columns

HEADER_FONT = Font(bold=True, color="FFFFFF")
HEADER_FILL = PatternFill("solid", fgColor="4472C4")
THIN_SIDE = Side(style="thin", color="D9D9D9")
THIN_BORDER = Border(left=THIN_SIDE, right=THIN_SIDE, top=THIN_SIDE, bottom=THIN_SIDE)
ID_ROW = 2
HEADER_ROW = 1

# Ko'p tilli statik yorliqlar. Frontend qaysi tilda bo'lsa, o'sha `lang`
# qiymatini yuboradi (query param). Noma'lum til yoki lang berilmasa -> uz.
LABELS = {
    "uz": {
        "plant_id": "Plant ID (tizim uchun, o'chirmang)",
        "plant_name": "O'simlik nomi",
        "sheet_title": "Reja",
        "no_plants": "(o'simliklar topilmadi)",
        "no_plants_export": "(rejaga o'simlik kiritilmagan)",
    },
    "ru": {
        "plant_id": "ID растения (для системы, не удаляйте)",
        "plant_name": "Название растения",
        "sheet_title": "План",
        "no_plants": "(растения не найдены)",
        "no_plants_export": "(в план не добавлены растения)",
    },
}


def _label(key, lang):
    return LABELS.get(lang, LABELS["uz"]).get(key, LABELS["uz"][key])


def _column_title(column, lang):
    return getattr(column, f"title_{lang}", column.title_uz) or column.title_uz


def _plant_name(plant, lang):
    """
    O'simlik nomini saytda tanlangan tilda qaytaradi.

    Plant modelida `name_uz` / `name_ru` bo'lishi ham, faqat `name` bo'lishi
    ham mumkin bo'lgani uchun ikkalasiga ham mos ishlaydi:
        name_{lang} -> name_uz -> name
    """
    value = getattr(plant, f"name_{lang}", None)
    if value:
        return value

    value = getattr(plant, "name_uz", None)
    if value:
        return value

    return getattr(plant, "name", "") or ""


# Guruh yorlig'i sifatida ishlatiladigan parent title uchun maksimal uzunlik.
# Undan uzun sarlavhalar — bu statement'ning tavsifiy (uzun) sarlavha
# ustunlari (masalan "...2026-yilda tayyorlanadigan urug'larni...") bo'lib,
# ular har bir ustunda takrorlanadi va ajratuvchi ma'lumot bermaydi.
MAX_LABEL_PART_LEN = 40


def _full_column_label(column, lang, max_ancestors=2):
    """
    Ustun uchun ajratuvchi (parent ierarxiyasi + leaf) sarlavha.

    Ko'p statement'larda leaf ustun nomi bir xil ("Reja") bo'ladi — masalan
    Daromad'da 12 ta oy ustuni ham "Reja" deb nomlangan. Faqat leaf title'ni
    header qilib qo'ysak, foydalanuvchi qaysi ustun qaysi oyga tegishli
    ekanini ajrata olmaydi.

    Shu sababli leaf'ga eng yaqin qisqa guruh yorliqlarini (chorak, oy,
    kategoriya) qo'shamiz — masalan "1-chorak - Yanvar - Reja". Uzun tavsifiy
    parent sarlavhalar (statement header ustunlari) tashlab yuboriladi, chunki
    ular ajratuvchi emas va sarlavhani o'qib bo'lmas darajada uzun qiladi.
    """
    ancestors = []
    node = getattr(column, "parent", None)
    while node is not None:
        title = _column_title(node, lang)
        if title and len(title) <= MAX_LABEL_PART_LEN:
            ancestors.append(title)
        node = getattr(node, "parent", None)
    # leaf'ga eng yaqin `max_ancestors` ta qisqa yorliqni olamiz
    nearest = ancestors[:max_ancestors]
    nearest.reverse()
    nearest.append(_column_title(column, lang))
    return " - ".join(part for part in nearest if part)


def _autosize_header_widths(labels, min_width=14, max_width=32):
    """Sarlavha matni uzunligiga qarab taxminiy ustun kengligini hisoblaydi."""
    widths = []
    for label in labels:
        length = len(label or "")
        widths.append(max(min_width, min(max_width, int(length * 1.05) + 2)))
    return widths


def _style_header_cell(cell):
    cell.font = HEADER_FONT
    cell.fill = HEADER_FILL
    cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
    cell.border = THIN_BORDER


def build_plan_template(statement, lang="uz"):
    """
    Statement uchun to'ldirish shablonini (.xlsx, openpyxl Workbook) quradi.

    - use_plant_plan=False bo'lsa: 1 ta ma'lumot qatori (3-qator), har bir
      ustun uchun bitta katak.
    - use_plant_plan=True bo'lsa: har bir leaf Plant uchun bitta qator
      (3-qatordan boshlab). A-ustunida Plant ID saqlanadi, lekin bu ustun
      YASHIRILGAN — foydalanuvchi faqat B-ustundagi (tanlangan tildagi)
      o'simlik nomini ko'radi.
    """
    columns = list(get_plan_input_columns(statement))
    if not columns:
        raise ValueError(
            "Bu statement uchun to'ldiriladigan ustunlar topilmadi "
            "(status_plan sozlanmagan bo'lishi mumkin)."
        )

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = _label("sheet_title", lang)

    is_plant_plan = bool(getattr(statement, "use_plant_plan", False))
    col_offset = 3 if is_plant_plan else 1

    if is_plant_plan:
        id_cell = ws.cell(row=HEADER_ROW, column=1, value=_label("plant_id", lang))
        _style_header_cell(id_cell)

        name_cell = ws.cell(row=HEADER_ROW, column=2, value=_label("plant_name", lang))
        _style_header_cell(name_cell)

    column_labels = [_full_column_label(column, lang) for column in columns]

    for i, (column, label) in enumerate(zip(columns, column_labels)):
        cell = ws.cell(row=HEADER_ROW, column=col_offset + i, value=label)
        _style_header_cell(cell)
        ws.cell(row=ID_ROW, column=col_offset + i, value=str(column.id))

    ws.row_dimensions[ID_ROW].hidden = True
    ws.row_dimensions[HEADER_ROW].height = 42

    if is_plant_plan:
        plants = (
            Plant.objects.filter(statement=statement, children__isnull=True)
            .order_by("name")
        )
        row = 3
        for plant in plants:
            id_cell = ws.cell(row=row, column=1, value=str(plant.id))
            id_cell.border = THIN_BORDER

            name_cell = ws.cell(row=row, column=2, value=_plant_name(plant, lang))
            name_cell.border = THIN_BORDER
            row += 1
        if row == 3:
            # Statement'ga bog'liq leaf plant topilmadi — bo'sh qator qoldiramiz
            ws.cell(row=3, column=2, value=_label("no_plants", lang))
    else:
        # Yagona ma'lumot qatori uchun joy
        ws.cell(row=3, column=1, value=None)

    total_cols = col_offset + len(columns)

    if is_plant_plan:
        # A-ustun (Plant ID) foydalanuvchiga ko'rinmaydi, lekin import uchun
        # faylda saqlanadi.
        ws.column_dimensions["A"].hidden = True
        ws.column_dimensions["A"].width = 8
        ws.column_dimensions["B"].width = 30

    data_widths = _autosize_header_widths(column_labels)
    for i, width in enumerate(data_widths):
        ws.column_dimensions[get_column_letter(col_offset + i)].width = width

    ws.freeze_panes = get_column_letter(col_offset) + "3"
    ws.sheet_view.showGridLines = True
    return wb


def _ordering_sort_key(column):
    """`ordering` maydonini ("1.3.2.1") raqamli tuple sifatida saralaydi —
    string sort ko'p raqamli segmentlarni ("1.10" vs "1.3") noto'g'ri
    tartiblamasligi uchun."""
    parts = []
    for seg in str(getattr(column, "ordering", "") or "0").split("."):
        try:
            parts.append(int(seg))
        except ValueError:
            parts.append(0)
    return parts


def _export_columns(statement):
    """Eksport uchun ustunlar — foydalanuvchi kiritadigan (oylik) VA hisoblangan
    (chorak / yillik / jami) barcha status_plan leaf ustunlari, reja tartibida."""
    columns = list(
        StatementColumn.objects.filter(
            statement=statement,
            parent__isnull=False,
            status_plan__isnull=False,
        ).select_related(
            "parent", "parent__parent", "parent__parent__parent",
            "parent__parent__parent__parent",
        )
    )
    columns.sort(key=_ordering_sort_key)
    return columns


def build_plan_export(plan, lang="uz"):
    """
    Yaratilgan Plan'ni to'ldirilgan (o'qish uchun) .xlsx sifatida quradi.

    Layout:
      - use_plant_plan=False -> 1-qator sarlavha, 2-qator qiymatlar.
      - use_plant_plan=True  -> har bir plant uchun bitta qator (A: tanlangan
        tildagi o'simlik nomi), ustunlarda shu plant uchun qiymatlar.

    Ustunlar to'liq breakdown: foydalanuvchi kiritgan oylik qiymatlar hamda
    ular asosida hisoblangan chorak / yillik / jami ustunlari ham ko'rsatiladi.
    Jadval sarlavha uzunligiga qarab avtomatik kengaytiriladi, grid chiziqlari,
    autofilter va freeze panes bilan chiroyli ko'rinishda tayyorlanadi.
    """
    statement = plan.statement
    columns = _export_columns(statement)
    if not columns:
        raise ValueError(
            "Bu reja uchun ustunlar topilmadi (status_plan sozlanmagan bo'lishi mumkin)."
        )

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = _label("sheet_title", lang)

    is_plant_plan = bool(getattr(statement, "use_plant_plan", False))
    col_offset = 2 if is_plant_plan else 1

    if is_plant_plan:
        name_header = ws.cell(row=HEADER_ROW, column=1, value=_label("plant_name", lang))
        _style_header_cell(name_header)

    column_labels = [_full_column_label(column, lang) for column in columns]

    for i, (column, label) in enumerate(zip(columns, column_labels)):
        cell = ws.cell(row=HEADER_ROW, column=col_offset + i, value=label)
        _style_header_cell(cell)

    ws.row_dimensions[HEADER_ROW].height = 42

    last_row = HEADER_ROW

    if is_plant_plan:
        row = 2
        plant_plans = (
            plan.plant_plans
            .select_related("plant")
            .prefetch_related("items")
            .all()
        )
        for pp in plant_plans:
            name_cell = ws.cell(row=row, column=1, value=_plant_name(pp.plant, lang))
            name_cell.border = THIN_BORDER

            values = {item.column_id: item.plan_value for item in pp.items.all()}
            for i, column in enumerate(columns):
                value = values.get(column.id)
                cell = ws.cell(row=row, column=col_offset + i)
                cell.border = THIN_BORDER
                if value is not None:
                    cell.value = float(value)
                    cell.number_format = "#,##0.00"
            row += 1
        if row == 2:
            ws.cell(row=2, column=1, value=_label("no_plants_export", lang))
            row = 3
        last_row = row - 1
    else:
        values = {item.column_id: item.plan_value for item in plan.plan_items.all()}
        for i, column in enumerate(columns):
            value = values.get(column.id)
            cell = ws.cell(row=2, column=col_offset + i)
            cell.border = THIN_BORDER
            if value is not None:
                cell.value = float(value)
                cell.number_format = "#,##0.00"
        last_row = 2

    total_cols = col_offset + len(columns) - 1

    if is_plant_plan:
        ws.column_dimensions["A"].width = 30

    data_widths = _autosize_header_widths(column_labels)
    for i, width in enumerate(data_widths):
        ws.column_dimensions[get_column_letter(col_offset + i)].width = width

    # Chiroyli grid ko'rinish uchun autofilter va freeze panes
    last_col_letter = get_column_letter(total_cols)
    ws.auto_filter.ref = f"A{HEADER_ROW}:{last_col_letter}{last_row}"
    ws.freeze_panes = get_column_letter(col_offset) + "2"
    ws.sheet_view.showGridLines = True

    return wb


def parse_plan_excel(file_obj, statement):
    """
    Yuklangan Excel faylni o'qib, PlanSerializer ga mos payload qaytaradi:
        - use_plant_plan=False -> {"plan_items": [{"column": <uuid>, "plan_value": <num>}, ...]}
        - use_plant_plan=True  -> {"plants": [{"plant": <uuid>, "items": [...]}, ...]}

    Bo'sh kataklar e'tiborga olinmaydi (0 sifatida emas — umuman yuborilmaydi).
    Plant ID ustuni (A) yashirilgan bo'lsa ham qiymati faylda saqlanib qoladi,
    shuning uchun bu yerda hidden holatidan qat'iy nazar to'g'ri o'qiladi.
    """
    check_xlsx_zip(file_obj)  # audit R1: zip bomb
    try:
        wb = openpyxl.load_workbook(file_obj, data_only=True)
    except Exception as exc:  # noqa: BLE001
        raise ValueError(f"Fayl noto'g'ri formatda: {exc}")

    ws = wb.active
    is_plant_plan = bool(getattr(statement, "use_plant_plan", False))
    col_offset = 3 if is_plant_plan else 1

    id_row_cells = ws[ID_ROW] if ws.max_row >= ID_ROW else []
    column_ids = {}
    for cell in id_row_cells:
        if cell.column < col_offset:
            continue
        if cell.value:
            column_ids[cell.column] = str(cell.value).strip()

    if not column_ids:
        raise ValueError(
            "Shablon noto'g'ri: ustun ID'lari topilmadi. "
            "Iltimos, ushbu statement uchun tizimdan yuklab olingan shablondan foydalaning."
        )

    def _row_items(row_cells):
        items = []
        for cell in row_cells:
            if cell.column < col_offset:
                continue
            col_id = column_ids.get(cell.column)
            if not col_id:
                continue
            value = cell.value
            if value in (None, ""):
                continue
            items.append({"column": col_id, "plan_value": value})
        return items

    if is_plant_plan:
        plants_payload = []
        for row_cells in ws.iter_rows(min_row=3):
            plant_id = row_cells[0].value if row_cells else None
            if not plant_id:
                continue
            items = _row_items(row_cells)
            if items:
                plants_payload.append({"plant": str(plant_id).strip(), "items": items})

        if not plants_payload:
            raise ValueError("Faylda hech qanday o'simlik uchun qiymat topilmadi.")

        return {"plants": plants_payload}

    if ws.max_row < 3:
        raise ValueError("Faylda ma'lumot qatori (3-qator) topilmadi.")

    items = _row_items(ws[3])
    if not items:
        raise ValueError("Faylda hech qanday qiymat topilmadi.")

    return {"plan_items": items}    