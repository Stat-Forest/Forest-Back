from django.contrib import admin, messages
from django.http import HttpResponse
import csv
import io
from docx import Document
from reportlab.pdfgen import canvas

from apps.general.models.land_document import LandDocument


@admin.register(LandDocument)
class LandDocumentAdmin(admin.ModelAdmin):
    list_display = ("name", "document_type", "organization", "date", "created_at")
    list_filter = ("document_type", "organization", "date")
    search_fields = ("name", "organization__name")
    date_hierarchy = "date"
    actions = ["export_as_excel", "export_as_word", "export_as_pdf"]

    # ✅ Helper function: filter only uploaded files
    def _get_queryset_with_files(self, queryset):
        """Return only documents that have an uploaded file."""
        return queryset.exclude(file="").exclude(file__isnull=True)

    # ✅ Excel Export (CSV)
    def export_as_excel(self, request, queryset):
        queryset = self._get_queryset_with_files(queryset)
        if not queryset.exists():
            self.message_user(request, "No documents with attached files found!", level=messages.WARNING)
            return

        response = HttpResponse(content_type='text/csv')
        response['Content-Disposition'] = 'attachment; filename="land_documents.csv"'
        writer = csv.writer(response)
        writer.writerow(["Name", "Date", "Type", "Organization", "Comment", "File URL"])

        for obj in queryset:
            file_url = obj.file.url if obj.file else ""
            writer.writerow([
                obj.name,
                obj.date,
                obj.get_document_type_display(),
                obj.organization.name,
                obj.comment or "",
                file_url
            ])
        return response

    export_as_excel.short_description = "📊 Export (only with files) to Excel (CSV)"

    # ✅ Word Export
    def export_as_word(self, request, queryset):
        queryset = self._get_queryset_with_files(queryset)
        if not queryset.exists():
            self.message_user(request, "No documents with attached files found!", level=messages.WARNING)
            return

        document = Document()
        document.add_heading('Land Documents (With Files)', 0)

        for obj in queryset:
            document.add_paragraph(f"Name: {obj.name}")
            document.add_paragraph(f"Date: {obj.date}")
            document.add_paragraph(f"Type: {obj.get_document_type_display()}")
            document.add_paragraph(f"Organization: {obj.organization.name}")
            if obj.comment:
                document.add_paragraph(f"Comment: {obj.comment}")
            if obj.file:
                document.add_paragraph(f"File URL: {obj.file.url}")
            document.add_paragraph("-" * 50)

        buffer = io.BytesIO()
        document.save(buffer)
        buffer.seek(0)

        response = HttpResponse(
            buffer.getvalue(),
            content_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document'
        )
        response['Content-Disposition'] = 'attachment; filename="land_documents_with_files.docx"'
        return response

    export_as_word.short_description = "📝 Export (only with files) to Word"

    # ✅ PDF Export
    def export_as_pdf(self, request, queryset):
        queryset = self._get_queryset_with_files(queryset)
        if not queryset.exists():
            self.message_user(request, "No documents with attached files found!", level=messages.WARNING)
            return

        response = HttpResponse(content_type='application/pdf')
        response['Content-Disposition'] = 'attachment; filename="land_documents_with_files.pdf"'
        pdf = canvas.Canvas(response)

        y = 800
        pdf.setFont("Helvetica", 12)
        pdf.drawString(200, y + 20, "Land Documents (With Files)")
        y -= 40

        for obj in queryset:
            lines = [
                f"Name: {obj.name}",
                f"Date: {obj.date}",
                f"Type: {obj.get_document_type_display()}",
                f"Organization: {obj.organization.name}",
                f"Comment: {obj.comment or ''}",
                f"File URL: {obj.file.url if obj.file else ''}",
                "-" * 80,
            ]
            for line in lines:
                pdf.drawString(50, y, line)
                y -= 20
                if y < 50:
                    pdf.showPage()
                    pdf.setFont("Helvetica", 12)
                    y = 800
        pdf.save()
        return response

    export_as_pdf.short_description = "📄 Export (only with files) to PDF"
