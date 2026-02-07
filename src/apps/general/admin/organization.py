from django.contrib import admin, messages
from modeltranslation.admin import TranslationAdmin
from django.utils.translation import gettext_lazy as _
from django.utils import timezone
from datetime import timedelta
from django.db.models import Q, Max, Subquery, OuterRef
from django.http import HttpResponse
import csv
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

from apps.general.models import Organization
from apps.logs.mixins import AdminHistoryMixin
from apps.logs.models import HistoryActionJournal


class LoginStatusFilter(admin.SimpleListFilter):
    title = _('Login holati')
    parameter_name = 'login_status'

    def lookups(self, request, model_admin):
        return (
            ("never", _("Hech qachon sitega kirmagan")),
            ("7d", _("Oxirgi 7 kun ichida kirmagan")),
            ("1m", _("Oxirgi 1 oy ichida kirmagan")),
            ("3m", _("Oxirgi 3 oy ichida kirmagan")),
        )

    def queryset(self, request, queryset):
        if self.value() is None:
            return queryset

        now = timezone.now()

        if self.value() == "never":
            # Hech qachon login qilmagan organizatsiyalar
            org_ids_with_login = HistoryActionJournal.objects.filter(
                action='login',
                user__organization__isnull=False
            ).values_list('user__organization_id', flat=True).distinct()

            return queryset.exclude(id__in=org_ids_with_login)

        # Vaqt chegaralarini aniqlash
        if self.value() == "7d":
            date_threshold = now - timedelta(days=7)
        elif self.value() == "1m":
            date_threshold = now - timedelta(days=30)
        elif self.value() == "3m":
            date_threshold = now - timedelta(days=90)
        else:
            return queryset

        # Har bir organizatsiya uchun eng oxirgi login vaqtini topish
        latest_login_per_org = HistoryActionJournal.objects.filter(
            user__organization=OuterRef('pk'),
            action='login'
        ).order_by('-created_at').values('created_at')[:1]

        # Oxirgi login vaqti threshold dan oldingi yoki hech qachon login qilmagan organizatsiyalar
        org_ids = queryset.annotate(
            last_login_time=Subquery(latest_login_per_org)
        ).filter(
            Q(last_login_time__lt=date_threshold) | Q(last_login_time__isnull=True)
        ).values_list('id', flat=True)

        return queryset.filter(id__in=org_ids)


@admin.register(Organization)
class OrganizationAdmin(AdminHistoryMixin, TranslationAdmin):
    list_display = ('name_uz', 'region', 'stir', 'get_last_login', 'get_login_status_display')
    list_filter = ('region', LoginStatusFilter)
    search_fields = ('name', 'name_uz', 'name_ru', 'stir')
    actions = ['export_to_csv', 'export_to_excel']

    def get_last_login(self, obj):
        """Organizatsiya userlarining oxirgi login vaqti"""
        last_login = HistoryActionJournal.objects.filter(
            user__organization=obj,
            action='login'
        ).order_by('-created_at').first()

        if last_login:
            return last_login.created_at.strftime('%Y-%m-%d %H:%M:%S')
        return "Hech qachon"

    get_last_login.short_description = _("Oxirgi login")

    def get_login_status_display(self, obj):
        """Login holatini rangli ko'rsatish"""
        last_login = HistoryActionJournal.objects.filter(
            user__organization=obj,
            action='login'
        ).order_by('-created_at').first()

        if not last_login:
            return "❌ Hech qachon"

        now = timezone.now()
        days_diff = (now - last_login.created_at).days

        if days_diff <= 7:
            return "✅ Faol"
        elif days_diff <= 30:
            return "⚠️ 1 oy ichida kirmagan"
        elif days_diff <= 90:
            return "⚠️ 3 oy ichida kirmagan"
        else:
            return "❌ 3+ oy ichida kirmagan"

    get_login_status_display.short_description = _("Login holati")

    def get_login_status(self, obj):
        """Login holatini aniqlash"""
        last_login = HistoryActionJournal.objects.filter(
            user__organization=obj,
            action='login'
        ).order_by('-created_at').first()

        if not last_login:
            return "Hech qachon sitega kirmagan"

        now = timezone.now()
        days_diff = (now - last_login.created_at).days

        if days_diff <= 7:
            return "Faol"
        elif days_diff <= 30:
            return "Oxirgi 1 oy ichida kirmagan"
        elif days_diff <= 90:
            return "Oxirgi 3 oy ichida kirmagan"
        else:
            return "Oxirgi 3 oydan ko'proq vaqt ichida kirmagan"

    def export_to_csv(self, request, queryset):
        """CSV formatda export qilish"""
        response = HttpResponse(content_type='text/csv; charset=utf-8')
        response['Content-Disposition'] = 'attachment; filename="organizations.csv"'
        response.write('\ufeff'.encode('utf-8'))  # BOM for Excel UTF-8 support

        writer = csv.writer(response)

        # Header
        writer.writerow([
            '№',
            'Tashkilot nomi',
            'Region',
            'STIR',
            'Oxirgi login',
            'Login holati'
        ])

        # Ma'lumotlar
        for index, org in enumerate(queryset, start=1):
            name_uz = org.name_uz if hasattr(org, 'name_uz') and org.name_uz else (org.name if org.name else '')
            region = str(org.region) if org.region else ''
            stir = str(org.stir) if hasattr(org, 'stir') and org.stir else ''
            last_login = self.get_last_login(org)
            login_status = self.get_login_status(org)

            writer.writerow([
                index,
                name_uz,
                region,
                stir,
                last_login,
                login_status
            ])

        return response

    export_to_csv.short_description = _("CSV formatda yuklab olish")

    def export_to_excel(self, request, queryset):
        """Excel formatda export qilish"""
        response = HttpResponse(
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        response['Content-Disposition'] = 'attachment; filename="tashkilotlar.xlsx"'

        wb = Workbook()
        ws = wb.active
        ws.title = "Tashkilotlar"

        # Stil uchun ranglar
        header_fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
        header_font = Font(bold=True, color="FFFFFF", size=12)
        header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

        # Border
        thin_border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )

        # Data alignment
        data_alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        center_alignment = Alignment(horizontal="center", vertical="center")

        # Headerlar
        headers = [
            '№',
            'Tashkilot nomi',
            'Region',
            'STIR',
            'Oxirgi login',
            'Login holati'
        ]

        # Header qo'shish
        for col_num, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col_num)
            cell.value = header
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = header_alignment
            cell.border = thin_border

        # Ma'lumotlar qo'shish
        for row_num, org in enumerate(queryset, start=2):
            # №
            cell = ws.cell(row=row_num, column=1)
            cell.value = row_num - 1
            cell.alignment = center_alignment
            cell.border = thin_border

            # Tashkilot nomi
            name_uz = org.name_uz if hasattr(org, 'name_uz') and org.name_uz else (org.name if org.name else '')
            cell = ws.cell(row=row_num, column=2)
            cell.value = name_uz
            cell.alignment = data_alignment
            cell.border = thin_border

            # Region
            cell = ws.cell(row=row_num, column=3)
            cell.value = str(org.region) if org.region else ''
            cell.alignment = data_alignment
            cell.border = thin_border

            # STIR
            cell = ws.cell(row=row_num, column=4)
            cell.value = str(org.stir) if hasattr(org, 'stir') and org.stir else ''
            cell.alignment = center_alignment
            cell.border = thin_border

            # Oxirgi login
            cell = ws.cell(row=row_num, column=5)
            cell.value = self.get_last_login(org)
            cell.alignment = center_alignment
            cell.border = thin_border

            # Login holati
            cell = ws.cell(row=row_num, column=6)
            cell.value = self.get_login_status(org)
            cell.alignment = center_alignment
            cell.border = thin_border

        # Ustunlar kengligini sozlash
        ws.column_dimensions['A'].width = 8   # №
        ws.column_dimensions['B'].width = 70  # Tashkilot nomi
        ws.column_dimensions['C'].width = 20  # Region
        ws.column_dimensions['D'].width = 15  # STIR
        ws.column_dimensions['E'].width = 22  # Oxirgi login
        ws.column_dimensions['F'].width = 30  # Login holati

        # Birinchi qatorni (header) balandligini sozlash
        ws.row_dimensions[1].height = 30

        # Barcha qatorlarni avtomatik balandlikka sozlash
        for row in range(2, ws.max_row + 1):
            ws.row_dimensions[row].height = None  # Auto-height

        # Freeze panes - header qatorini doimo ko'rinishda qoldirish
        ws.freeze_panes = 'A2'

        wb.save(response)
        return response

    export_to_excel.short_description = _("Excel formatda yuklab olish")

    def save_model(self, request, obj, form, change):
        if Organization.objects.exclude(pk=obj.pk).filter(name=obj.name).exists():
            self.message_user(request, _("An organization with this name already exists."), level=messages.ERROR)
            return
        super().save_model(request, obj, form, change)