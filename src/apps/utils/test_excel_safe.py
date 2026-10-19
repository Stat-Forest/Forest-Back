from django.test import SimpleTestCase
from openpyxl import Workbook

from apps.utils.excel_safe import append_safe, set_safe


class ExcelSafeCellTest(SimpleTestCase):
    """Audit B6: foydalanuvchi matni Excel'da formula sifatida emas, matn sifatida saqlanadi."""

    def test_formula_like_text_is_stored_as_string(self):
        ws = Workbook().active
        cell = set_safe(ws.cell(row=1, column=1), '=HYPERLINK("https://evil","x")')
        self.assertEqual(cell.data_type, "s")
        self.assertEqual(cell.value, '=HYPERLINK("https://evil","x")')

    def test_numbers_and_plain_text_unchanged(self):
        ws = Workbook().active
        self.assertEqual(set_safe(ws.cell(row=1, column=1), 12.5).data_type, "n")
        self.assertEqual(set_safe(ws.cell(row=1, column=2), "Salom").data_type, "s")
        self.assertIsNone(set_safe(ws.cell(row=1, column=3), None).value)

    def test_append_safe_neutralises_every_string_cell(self):
        ws = Workbook().active
        append_safe(ws, ["=1+1", "+cmd", 5])
        row = ws[ws.max_row]
        self.assertEqual([c.data_type for c in row], ["s", "s", "n"])
