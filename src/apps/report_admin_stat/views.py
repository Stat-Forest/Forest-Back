# apps/report_admin_stat/views.py
import csv
from io import BytesIO
from django.http import HttpResponse
from django.utils import timezone
from rest_framework.views import APIView
from rest_framework.response import Response
from rest_framework.permissions import IsAuthenticated

from .services import StatementReportStatService
from .serializers import StatementStatResponseSerializer

from drf_spectacular.utils import extend_schema


@extend_schema(tags=["Dashboards"])
class StatementReportStatView(APIView):
    permission_classes = [IsAuthenticated]

    def get(self, request):
        service = StatementReportStatService(request.user)
        statistics = service.get_statistics()
        totals = service.get_totals(statistics)

        serializer = StatementStatResponseSerializer({
            "data":   statistics,
            "totals": totals,
        })
        return Response(serializer.data)


@extend_schema(
    tags=["Stats Export"],
)
class StatementReportExportView(APIView):
    permission_classes = [IsAuthenticated]

    def get(self, request):
        export_format = request.query_params.get('format', 'excel')
        service = StatementReportStatService(request.user)
        statistics = service.get_statistics()
        totals = service.get_totals(statistics)

        if export_format == 'csv':
            return self._export_csv(statistics, totals)
        return self._export_excel(statistics, totals)

    # ================= FLAT DATA =================
    def _flatten(self, statistics):
        rows = []
        for parent in statistics:
            rows.append({
                "name":      parent["name"],
                "is_parent": True,
                "saved":     parent["saved"],
                "sent":      parent["sent"],
                "cancelled": parent["cancelled"],
                "accepted":  parent["accepted"],
                "total":     parent["total"],
            })
            for child in parent.get("children", []):
                rows.append({
                    "name":      f"  {child['name']}",
                    "is_parent": False,
                    "saved":     child["saved"],
                    "sent":      child["sent"],
                    "cancelled": child["cancelled"],
                    "accepted":  child["accepted"],
                    "total":     child["total"],
                })
        return rows

    # ================= CSV =================
    def _export_csv(self, statistics, totals):
        response = HttpResponse(content_type='text/csv; charset=utf-8')
        timestamp = timezone.now().strftime('%Y%m%d_%H%M%S')
        response['Content-Disposition'] = f'attachment; filename="statement_statistics_{timestamp}.csv"'
        response.write('\ufeff')

        writer = csv.writer(response)
        writer.writerow([
            'Hisobot formasi', 'Saqlangan', 'Yuborilgan',
            'Bekor qilingan', 'Qabul qilingan', 'Jami'
        ])

        for row in self._flatten(statistics):
            writer.writerow([
                row["name"], row["saved"], row["sent"],
                row["cancelled"], row["accepted"], row["total"],
            ])

        writer.writerow([])
        writer.writerow([
            'JAMI:', totals["saved"], totals["sent"],
            totals["cancelled"], totals["accepted"], totals["total"],
        ])

        return response

    # ================= EXCEL =================
    def _export_excel(self, statistics, totals):
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        except ImportError:
            return HttpResponse("openpyxl o'rnatilmagan: pip install openpyxl", status=500)

        wb = Workbook()
        ws = wb.active
        ws.title = "Statistika"

        HEADER_FILL = PatternFill(start_color="667eea", end_color="667eea", fill_type="solid")
        PARENT_FILL = PatternFill(start_color="EEF2FF", end_color="EEF2FF", fill_type="solid")
        CHILD_FILL  = PatternFill(start_color="F8F9FA", end_color="F8F9FA", fill_type="solid")
        TOTAL_FILL  = PatternFill(start_color="FEF9C3", end_color="FEF9C3", fill_type="solid")
        THIN_BORDER = Border(
            left=Side(style='thin'), right=Side(style='thin'),
            top=Side(style='thin'),  bottom=Side(style='thin')
        )

        headers = [
            'Hisobot formasi', 'Saqlangan', 'Yuborilgan',
            'Bekor qilingan', 'Qabul qilingan', 'Jami'
        ]

        for col, h in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=h)
            cell.fill      = HEADER_FILL
            cell.font      = Font(bold=True, color="FFFFFF", size=12)
            cell.alignment = Alignment(horizontal='center', vertical='center')
            cell.border    = THIN_BORDER
        ws.row_dimensions[1].height = 30

        row_idx = 2
        for row in self._flatten(statistics):
            values = [
                row["name"], row["saved"], row["sent"],
                row["cancelled"], row["accepted"], row["total"],
            ]
            for col, val in enumerate(values, 1):
                cell = ws.cell(row=row_idx, column=col, value=val)
                cell.border    = THIN_BORDER
                cell.alignment = Alignment(
                    horizontal='left' if col == 1 else 'center',
                    vertical='center'
                )
                if row["is_parent"]:
                    cell.font = Font(bold=True)
                    cell.fill = PARENT_FILL
                else:
                    cell.fill = CHILD_FILL
            row_idx += 1

        total_row_idx = row_idx + 1
        total_values = [
            'JAMI:', totals["saved"], totals["sent"],
            totals["cancelled"], totals["accepted"], totals["total"],
        ]
        for col, val in enumerate(total_values, 1):
            cell = ws.cell(row=total_row_idx, column=col, value=val)
            cell.font      = Font(bold=True, size=12)
            cell.fill      = TOTAL_FILL
            cell.border    = THIN_BORDER
            cell.alignment = Alignment(
                horizontal='left' if col == 1 else 'center',
                vertical='center'
            )

        ws.column_dimensions['A'].width = 52
        for col in ['B', 'C', 'D', 'E', 'F']:
            ws.column_dimensions[col].width = 18

        buffer = BytesIO()
        wb.save(buffer)
        buffer.seek(0)

        timestamp = timezone.now().strftime('%Y%m%d_%H%M%S')
        response = HttpResponse(
            buffer,
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        response['Content-Disposition'] = f'attachment; filename="statement_statistics_{timestamp}.xlsx"'
        return response