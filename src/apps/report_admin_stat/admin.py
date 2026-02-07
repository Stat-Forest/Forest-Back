# apps/report_admin_stat/admin.py

from django.contrib import admin
from django.db.models import Count, Q
from django.http import HttpResponse
from django.utils import timezone
import csv
from apps.reports.models import Report
from .models import StatementStatistics


class StatementStatisticsAdmin(admin.ModelAdmin):
    """Parent statement-lar uchun hisobotlar statistikasi"""

    change_list_template = 'admin/reports/statement_statistics.html'

    def changelist_view(self, request, extra_context=None):
        # Export funksiyasi
        if request.GET.get('export') == 'csv':
            return self.export_to_csv(request)

        if request.GET.get('export') == 'excel':
            return self.export_to_excel(request)

        # Filter parametrlari
        status_filter = request.GET.get('status', '')
        specific_parent = request.GET.get('specific_parent', '')

        # FAQAT parent statement-larni olish (parent=None)
        if specific_parent:
            statements = StatementStatistics.objects.filter(id=specific_parent, parent__isnull=True)
        else:
            statements = StatementStatistics.objects.filter(parent__isnull=True)

        # Har bir parent statement uchun statistika
        statistics = []

        for statement in statements:
            # Parent statistikasini hisoblash
            child_statements = self.get_all_children(statement)
            all_statement_ids = [statement.id] + [s.id for s in child_statements]

            base_query = Report.objects.filter(
                statement_id__in=all_statement_ids,
                is_report=True
            )

            if status_filter:
                base_query = base_query.filter(status=status_filter)

            saved_count = base_query.filter(status=Report.StatusChoices.SAVED).count()
            sent_count = base_query.filter(status=Report.StatusChoices.SENT).count()
            cancelled_count = base_query.filter(status=Report.StatusChoices.CANCELLED).count()
            accepted_count = base_query.filter(status=Report.StatusChoices.ACCEPTED).count()
            total_count = base_query.count()

            # Parent ma'lumotlari
            parent_data = {
                'statement': statement,
                'statement_name_uz': getattr(statement, 'name_uz', statement.name),
                'is_parent': True,
                'child_count': len(child_statements),
                'saved': saved_count,
                'sent': sent_count,
                'cancelled': cancelled_count,
                'accepted': accepted_count,
                'total': total_count,
                'children': []
            }

            # Har bir child uchun alohida statistika
            for child in child_statements:
                child_query = Report.objects.filter(
                    statement_id=child.id,
                    is_report=True
                )

                if status_filter:
                    child_query = child_query.filter(status=status_filter)

                child_saved = child_query.filter(status=Report.StatusChoices.SAVED).count()
                child_sent = child_query.filter(status=Report.StatusChoices.SENT).count()
                child_cancelled = child_query.filter(status=Report.StatusChoices.CANCELLED).count()
                child_accepted = child_query.filter(status=Report.StatusChoices.ACCEPTED).count()
                child_total = child_query.count()

                parent_data['children'].append({
                    'statement': child,
                    'statement_name_uz': getattr(child, 'name_uz', child.name),
                    'is_parent': False,
                    'level': self.get_level(child),
                    'saved': child_saved,
                    'sent': child_sent,
                    'cancelled': child_cancelled,
                    'accepted': child_accepted,
                    'total': child_total,
                })

            statistics.append(parent_data)

        # Parent statement-larni dropdown uchun olish
        all_parent_statements = StatementStatistics.objects.filter(parent__isnull=True)

        # Umumiy statistikani hisoblash
        total_stats = {
            'total_reports': sum(s['total'] for s in statistics),
            'total_saved': sum(s['saved'] for s in statistics),
            'total_sent': sum(s['sent'] for s in statistics),
            'total_cancelled': sum(s['cancelled'] for s in statistics),
            'total_accepted': sum(s['accepted'] for s in statistics),
            'total_statements': len(statistics),
        }

        extra_context = extra_context or {}
        extra_context['statistics'] = statistics
        extra_context['status_choices'] = Report.StatusChoices.choices
        extra_context['current_status'] = status_filter
        extra_context['specific_parent'] = specific_parent
        extra_context['all_parent_statements'] = all_parent_statements
        extra_context['total_stats'] = total_stats
        extra_context['request'] = request  # request objektini qo'shish

        return super().changelist_view(request, extra_context=extra_context)

    def get_all_children(self, statement):
        """Rekursiv ravishda barcha child statement-larni olish"""
        children = []
        direct_children = StatementStatistics.objects.filter(parent=statement)

        for child in direct_children:
            children.append(child)
            children.extend(self.get_all_children(child))

        return children

    def get_level(self, statement):
        """Statement ning darajasini aniqlash"""
        level = 0
        current = statement
        while current.parent is not None:
            level += 1
            current = current.parent
        return level

    def has_add_permission(self, request):
        return False

    def has_delete_permission(self, request, obj=None):
        return False

    def has_change_permission(self, request, obj=None):
        return False

    def get_statistics_data(self, request):
        """Statistika ma'lumotlarini olish (export uchun)"""
        status_filter = request.GET.get('status', '')
        specific_parent = request.GET.get('specific_parent', '')

        if specific_parent:
            statements = StatementStatistics.objects.filter(id=specific_parent, parent__isnull=True)
        else:
            statements = StatementStatistics.objects.filter(parent__isnull=True)

        export_data = []

        for statement in statements:
            # Parent statistikasi
            child_statements = self.get_all_children(statement)
            all_statement_ids = [statement.id] + [s.id for s in child_statements]

            base_query = Report.objects.filter(
                statement_id__in=all_statement_ids,
                is_report=True
            )

            if status_filter:
                base_query = base_query.filter(status=status_filter)

            saved_count = base_query.filter(status=Report.StatusChoices.SAVED).count()
            sent_count = base_query.filter(status=Report.StatusChoices.SENT).count()
            cancelled_count = base_query.filter(status=Report.StatusChoices.CANCELLED).count()
            accepted_count = base_query.filter(status=Report.StatusChoices.ACCEPTED).count()
            total_count = base_query.count()

            export_data.append({
                'statement_name_uz': getattr(statement, 'name_uz', statement.name),
                'is_parent': True,
                'level': 0,
                'child_count': len(child_statements),
                'saved': saved_count,
                'sent': sent_count,
                'cancelled': cancelled_count,
                'accepted': accepted_count,
                'total': total_count,
            })

            # Child statistikalari
            for child in child_statements:
                child_query = Report.objects.filter(
                    statement_id=child.id,
                    is_report=True
                )

                if status_filter:
                    child_query = child_query.filter(status=status_filter)

                child_saved = child_query.filter(status=Report.StatusChoices.SAVED).count()
                child_sent = child_query.filter(status=Report.StatusChoices.SENT).count()
                child_cancelled = child_query.filter(status=Report.StatusChoices.CANCELLED).count()
                child_accepted = child_query.filter(status=Report.StatusChoices.ACCEPTED).count()
                child_total = child_query.count()

                export_data.append({
                    'statement_name_uz': getattr(child, 'name_uz', child.name),
                    'is_parent': False,
                    'level': self.get_level(child),
                    'child_count': 0,
                    'saved': child_saved,
                    'sent': child_sent,
                    'cancelled': child_cancelled,
                    'accepted': child_accepted,
                    'total': child_total,
                })

        return export_data

    def export_to_csv(self, request):
        """CSV formatda export qilish"""
        response = HttpResponse(content_type='text/csv; charset=utf-8')
        timestamp = timezone.now().strftime('%Y%m%d_%H%M%S')
        response['Content-Disposition'] = f'attachment; filename="statement_statistics_{timestamp}.csv"'

        response.write('\ufeff')
        writer = csv.writer(response)

        # Header
        writer.writerow([
            'Ҳисобот формаси',
            'Даража',
            'Сақланган',
            'Юборилган',
            'Бекор қилинган',
            'Қабул қилинган',
            'Жами'
        ])

        # Data
        statistics = self.get_statistics_data(request)
        for stat in statistics:
            indent = '  ' * stat['level']
            writer.writerow([
                indent + stat['statement_name_uz'],
                stat['level'],
                stat['saved'],
                stat['sent'],
                stat['cancelled'],
                stat['accepted'],
                stat['total'],
            ])

        # Jami qator
        writer.writerow([])
        parent_stats = [s for s in statistics if s['is_parent']]
        writer.writerow([
            'ЖАМИ:',
            '',
            sum(s['saved'] for s in parent_stats),
            sum(s['sent'] for s in parent_stats),
            sum(s['cancelled'] for s in parent_stats),
            sum(s['accepted'] for s in parent_stats),
            sum(s['total'] for s in parent_stats),
        ])

        return response

    def export_to_excel(self, request):
        """Excel formatda export qilish"""
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        except ImportError:
            return HttpResponse("openpyxl kutubxonasi o'rnatilmagan. pip install openpyxl", status=500)

        wb = Workbook()
        ws = wb.active
        ws.title = "Statistika"

        # Header styling
        header_fill = PatternFill(start_color="667eea", end_color="667eea", fill_type="solid")
        header_font = Font(bold=True, color="FFFFFF", size=12)
        border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )

        # Headers
        headers = [
            'Ҳисобот формаси',
            'Сақланган',
            'Юборилган',
            'Бекор қилинган',
            'Қабул қилинган',
            'Жами'
        ]

        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal='center', vertical='center')
            cell.border = border

        # Data
        statistics = self.get_statistics_data(request)
        row_idx = 2

        for stat in statistics:
            indent = '  ' * stat['level']
            ws.cell(row=row_idx, column=1, value=indent + stat['statement_name_uz'])
            ws.cell(row=row_idx, column=2, value=stat['saved'])
            ws.cell(row=row_idx, column=3, value=stat['sent'])
            ws.cell(row=row_idx, column=4, value=stat['cancelled'])
            ws.cell(row=row_idx, column=5, value=stat['accepted'])
            ws.cell(row=row_idx, column=6, value=stat['total'])

            # Styling
            for col in range(1, 7):
                cell = ws.cell(row=row_idx, column=col)
                cell.border = border
                cell.alignment = Alignment(horizontal='left' if col == 1 else 'center', vertical='center')

                if stat['is_parent']:
                    cell.font = Font(bold=True)
                    cell.fill = PatternFill(start_color="f0f4ff", end_color="f0f4ff", fill_type="solid")
                elif stat['level'] == 1:
                    cell.fill = PatternFill(start_color="f8f9fa", end_color="f8f9fa", fill_type="solid")

            row_idx += 1

        # Jami qator
        total_row = row_idx + 1
        parent_stats = [s for s in statistics if s['is_parent']]
        ws.cell(row=total_row, column=1, value='ЖАМИ:').font = Font(bold=True, size=12)
        ws.cell(row=total_row, column=2, value=sum(s['saved'] for s in parent_stats)).font = Font(bold=True)
        ws.cell(row=total_row, column=3, value=sum(s['sent'] for s in parent_stats)).font = Font(bold=True)
        ws.cell(row=total_row, column=4, value=sum(s['cancelled'] for s in parent_stats)).font = Font(bold=True)
        ws.cell(row=total_row, column=5, value=sum(s['accepted'] for s in parent_stats)).font = Font(bold=True)
        ws.cell(row=total_row, column=6, value=sum(s['total'] for s in parent_stats)).font = Font(bold=True)

        # Column widths
        ws.column_dimensions['A'].width = 50
        ws.column_dimensions['B'].width = 15
        ws.column_dimensions['C'].width = 15
        ws.column_dimensions['D'].width = 18
        ws.column_dimensions['E'].width = 18
        ws.column_dimensions['F'].width = 12

        # Response
        response = HttpResponse(
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        timestamp = timezone.now().strftime('%Y%m%d_%H%M%S')
        response['Content-Disposition'] = f'attachment; filename="statement_statistics_{timestamp}.xlsx"'

        wb.save(response)
        return response


# Proxy modelni registratsiya qilish
admin.site.register(StatementStatistics, StatementStatisticsAdmin)