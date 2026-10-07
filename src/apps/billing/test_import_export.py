"""Excel import (JSON qatorlar), Excel eksport, import shabloni, shartnoma fayllari."""
import io
import tempfile

import openpyxl
from django.core.cache import cache
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import override_settings
from rest_framework.test import APITestCase

from apps.billing._tests_common import BASE, auth, contract_payload, make_reference, make_user
from apps.billing.models import Contract, ContractFile, ContractHistory

URL = f"{BASE}/contracts/"


def import_row(ref, **overrides):
    row = {
        "contract": {
            "contract_date": "15.03.2025",
            "expired_date": 46096,  # Excel serial -> 15.03.2026
            "contract_number": "IMP-1",
            "xojalik_inn": ref.xojalik.inn,
            "bolim": "markaziy",
            "region_code": "3",
            "district_code": "5",
            "rent_area": "3,5",
            "rent_amount": "",
            "latitude": 40.71,
            "longitude": 72.31,
        },
        "renter": {"stir": "555555555", "full_name": "Import MCHJ", "is_legal_entity": True},
        "lands": [{"type": "чўл", "area": 2}, {"type": "Тоғ", "area": ""}],
        "activities": [{"type": "Ўрмон барпо қилиш", "area": "1.5"}],
        "privileges": [],
        "action": {"vacancy": "", "own_investment": "12"},
        "action_products": [{"product_export": "", "product_amount": "", "export_amount": ""}],
        "yearly_amounts": [{"year": 2025, "contract_amount": 100, "calculated_amount": 100, "paid_amount": 40, "penalty_amount": ""}],
    }
    row["contract"].update(overrides)
    return row


class ImportExcelTest(APITestCase):
    def setUp(self):
        cache.clear()
        self.ref = make_reference()
        self.client = auth(self.client, make_user("billing", "billing_user"))

    def test_import_valid_and_invalid_rows(self):
        rows = [import_row(self.ref), import_row(self.ref, xojalik_inn="000000000", contract_number="IMP-2")]
        resp = self.client.post(f"{URL}import-excel/", {"rows": rows}, format="json")
        self.assertEqual(resp.status_code, 200, resp.data)
        self.assertEqual(resp.data["saved"], [0])
        self.assertEqual(list(resp.data["errors"]), ["1"])
        self.assertIn("xojalik_inn", resp.data["errors"]["1"]["contract"])

        contract = Contract.objects.get(contract_number="IMP-1")
        self.assertEqual(contract.xojalik, self.ref.xojalik)
        self.assertEqual(contract.bolim, self.ref.bolim)
        self.assertEqual(contract.region, self.ref.region)
        self.assertEqual(contract.district, self.ref.district)
        self.assertEqual(str(contract.contract_date), "2025-03-15")
        self.assertEqual(str(contract.expired_date), "2026-03-15")
        self.assertEqual(str(contract.rent_area), "3.50")
        self.assertEqual(str(contract.rent_amount), "100.00")  # yillikdan
        self.assertEqual(contract.renter.stir, "555555555")
        self.assertTrue(contract.is_entity)
        self.assertEqual(contract.lands.count(), 1)  # bo'sh maydonli tur tashlab ketiladi
        self.assertEqual(contract.activities.get().type, self.ref.activity_type)
        self.assertEqual(contract.actions.get().products.count(), 0)
        self.assertEqual(contract.history.get().reason, ContractHistory.Reason.IMPORTED)
        self.assertEqual(contract.status, "active")

    def test_import_canceled_row_becomes_inactive(self):
        row = import_row(self.ref, canceled_date="2025-12-01", cancel_reason="Bekor")
        resp = self.client.post(f"{URL}import-excel/", {"rows": [row]}, format="json")
        self.assertEqual(resp.data["saved"], [0], resp.data)
        contract = Contract.objects.get()
        self.assertEqual(contract.status, "inactive")
        self.assertEqual(str(contract.canceled_date), "2025-12-01")

    def test_import_row_errors_are_sectioned(self):
        row = import_row(self.ref, contract_date="", rent_area="abc")
        row["lands"] = [{"type": "yo'q", "area": "1"}]
        resp = self.client.post(f"{URL}import-excel/", {"rows": [row]}, format="json")
        errors = resp.data["errors"]["0"]
        self.assertEqual(errors["contract"]["contract_date"], "To'ldirilishi shart")
        self.assertIn("rent_area", errors["contract"])
        self.assertEqual(Contract.objects.count(), 0)

        resp = self.client.post(f"{URL}import-excel/", {"rows": [import_row(self.ref, rent_area="2") | {"lands": [{"type": "yo'q", "area": "1"}]}]}, format="json")
        self.assertIn("yo'q", resp.data["errors"]["0"]["lands"])
        self.assertEqual(Contract.objects.count(), 0)  # qator atomik

    def test_import_duplicate_contract_number(self):
        self.client.post(f"{URL}import-excel/", {"rows": [import_row(self.ref)]}, format="json")
        resp = self.client.post(f"{URL}import-excel/", {"rows": [import_row(self.ref)]}, format="json")
        self.assertIn("contract_number", resp.data["errors"]["0"]["contract"])

    def test_import_limits_and_permissions(self):
        resp = self.client.post(f"{URL}import-excel/", {"rows": [{}] * 201}, format="json")
        self.assertEqual(resp.status_code, 400)
        reader = auth(self.client_class(), make_user("reader", "rahbariyat"))
        self.assertEqual(reader.post(f"{URL}import-excel/", {"rows": [import_row(self.ref)]}, format="json").status_code, 403)

    def test_import_template_download(self):
        resp = self.client.get(f"{URL}import-template/")
        self.assertEqual(resp.status_code, 200)
        self.assertIn("spreadsheetml", resp["Content-Type"])
        self.assertIn("shartnoma-import-namuna.xlsx", resp["Content-Disposition"])


class ExportExcelTest(APITestCase):
    def setUp(self):
        cache.clear()
        self.ref = make_reference()
        self.client = auth(self.client, make_user("billing", "billing_user"))
        self.client.post(URL, contract_payload(self.ref), format="json")
        resp = self.client.post(URL, contract_payload(self.ref, contract_number="SH-2", pinfl="22222222222222"), format="json")
        self.client.post(f"{URL}{resp.data['id']}/status/", {"status": "inactive", "cancel_reason": "x"}, format="json")

    def test_export_respects_filters_and_has_total_row(self):
        resp = self.client.get(f"{URL}export-excel/", {"status": "active"})
        self.assertEqual(resp.status_code, 200)
        self.assertIn("shartnomalar-", resp["Content-Disposition"])
        ws = openpyxl.load_workbook(io.BytesIO(resp.content))["1"]
        self.assertEqual(ws["C5"].value, self.ref.xojalik.name)
        self.assertEqual(ws["B5"].value, "123456789")
        self.assertEqual(ws["I5"].value, "SH-1")
        self.assertEqual(ws["J5"].value, "Жисмоний")
        self.assertEqual(ws["K5"].value, "12345678901234")
        self.assertEqual(ws["P5"].value, "Markaziy")
        self.assertEqual(ws["R5"].value, "3")
        self.assertEqual(float(ws["Y5"].value), 12.50)
        self.assertEqual(float(ws["Z5"].value), 7.50)   # faoliyat: ўрмон барпо қилиш
        self.assertEqual(float(ws["AS5"].value), 5.00)  # yer turi: Чўл
        self.assertEqual(float(ws["BB5"].value), 1.00)  # imtiyoz: доривор ... (3 йил)
        self.assertEqual(float(ws["BC5"].value), 1100.00)
        self.assertEqual(float(ws["BE5"].value), 500.00)
        self.assertEqual(float(ws["BF5"].value), 600.00)
        self.assertEqual(ws["BO5"].value, 40.7)
        self.assertEqual(float(ws["BV5"].value), 500.00)  # 2025 shartnoma summasi (BQ + 5)
        self.assertEqual(float(ws["CD5"].value), 600.00)  # 2026 hisoblangan (BX + 6)
        self.assertEqual(ws["C6"].value, "ЖАМИ")          # faqat 1 ta faol shartnoma
        self.assertEqual(ws["Y6"].value, "=SUM(Y5:Y5)")

        all_rows = openpyxl.load_workbook(io.BytesIO(self.client.get(f"{URL}export-excel/", {"status": "all"}).content))["1"]
        self.assertEqual(all_rows["C7"].value, "ЖАМИ")
        self.assertEqual(all_rows["BN6"].value, "x")


class ContractFilesTest(APITestCase):
    def setUp(self):
        cache.clear()
        self.ref = make_reference()
        self.user = make_user("billing", "billing_user")
        self.client = auth(self.client, self.user)
        self.contract_id = self.client.post(URL, contract_payload(self.ref), format="json").data["id"]
        self.media = tempfile.TemporaryDirectory()
        self.addCleanup(self.media.cleanup)

    def test_upload_list_delete(self):
        with override_settings(MEDIA_ROOT=self.media.name):
            pdf = SimpleUploadedFile("shartnoma.pdf", b"%PDF-1.4 test", content_type="application/pdf")
            resp = self.client.post(f"{BASE}/contract-files/", {"contract": self.contract_id, "file": pdf, "description": "Asl nusxa"}, format="multipart")
            self.assertEqual(resp.status_code, 201, resp.data)
            self.assertEqual(resp.data["file_name"], "shartnoma.pdf")
            self.assertIn("/media/documents/billing/contracts/", resp.data["file_url"])
            self.assertIn("?t=", resp.data["file_url"])
            self.assertEqual(resp.data["uploaded_by_name"], "T Billing")
            self.assertNotIn("file", resp.data)

            listed = self.client.get(f"{BASE}/contract-files/", {"contract": self.contract_id}).data
            self.assertEqual(listed["count"], 1)
            detail = self.client.get(f"{URL}{self.contract_id}/").data
            self.assertEqual(detail["files_count"], 1)
            self.assertEqual(detail["files"][0]["file_name"], "shartnoma.pdf")

            bad = SimpleUploadedFile("x.txt", b"hello", content_type="text/plain")
            self.assertEqual(self.client.post(f"{BASE}/contract-files/", {"contract": self.contract_id, "file": bad}, format="multipart").status_code, 400)

            self.assertEqual(self.client.delete(f"{BASE}/contract-files/{resp.data['id']}/").status_code, 204)
            self.assertEqual(ContractFile.objects.count(), 0)
        reasons = list(ContractHistory.objects.filter(contract_id=self.contract_id).order_by("created_at").values_list("reason", flat=True))
        self.assertEqual(reasons, ["created", "file_added", "file_deleted"])

    def test_reader_cannot_upload(self):
        reader = auth(self.client_class(), make_user("reader", "rahbariyat"))
        with override_settings(MEDIA_ROOT=self.media.name):
            pdf = SimpleUploadedFile("a.pdf", b"%PDF-1.4", content_type="application/pdf")
            self.assertEqual(reader.post(f"{BASE}/contract-files/", {"contract": self.contract_id, "file": pdf}, format="multipart").status_code, 403)
            self.assertEqual(reader.get(f"{BASE}/contract-files/", {"contract": self.contract_id}).status_code, 200)
