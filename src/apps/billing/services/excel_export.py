"""
Shartnomalar Excel hisoboti — `templates_xlsx/export_template.xlsx`
(billing'dagi billing_temp_back.xlsx, varaq "1", ma'lumot 5-qatordan) asosida,
billing'dagi ustun xaritasi + oxirida "ЖАМИ" qatori (=SUM formulalar).
Faoliyat/yer/imtiyoz turlari ustunlarga tur nomi bo'yicha (katta-kichik harf
va ortiqcha bo'sh joylarsiz) joylashtiriladi.
"""
from io import BytesIO
from pathlib import Path

from django.db.models import Prefetch, QuerySet
from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string, get_column_letter

from apps.billing.models import ContractActivity, ContractLand, ContractPrivilege, ContractYearlyAmount

TEMPLATE_PATH = Path(__file__).resolve().parents[1] / "templates_xlsx" / "export_template.xlsx"
SHEET_NAME = "1"
FIRST_DATA_ROW = 5
YEARS = list(range(2020, 2027))


def _norm(name: str) -> str:
    return " ".join(str(name or "").casefold().split())


def _columns(mapping: dict) -> dict:
    return {_norm(name): column_index_from_string(letter) for name, letter in mapping.items()}


ACTIVITY_COLUMNS = _columns({
    "ўрмон барпо қилиш": "Z",
    "қишлоқ хўжалиги экинларини экиш": "AA",
    "мевали боғлар барпо этиш": "AB",
    "ёнғоқ мевали боғлар барпо этиш (писта, бодом, ёнғоқ) фаолияти": "AC",
    "кўчат етиштириш фаолияти": "AD",
    "лимонарий ва иссиқхоналар ташкил этиш фаолияти": "AE",
    "пичанзорлардан фойдаланиш (пичан йиғиштириш)": "AF",
    "доривор ўсимликлар етиштириш фаолияти": "AG",
    "овчилик фаолияти (овчиларни қабул қилиш пункти)": "AH",
    "Тез ўсувчи дарахт плантацияларини барпо этиш учун": "AI",
    "Мевали ва манзарали кўчат етиштириш учун": "AJ",
    "Доривор ўсимликлар маданий плантацияларини барпо этиш учун": "AK",
    "Асаларичилик фаолияти учун": "AL",
    "Балиқчилик фаолияти учун": "AM",
    "Овчилик хўжалигини ташкил этиш учун": "AN",
    "Экологик туризм, рекрация, соғломлаштириш ҳамда дам олиш масканлари фаолияти учун (ўсимлик ва ҳайвонот дунёсига зарар келтирмаслик шарти билан)": "AO",
    "Бунда Бурчмулла, Оҳангарон, Китоб, Зомин, Бойсун, Узун, Жомбой, Бахмал давлат ўрмон хўжаликлари, Шаҳрисабз, Поп, Ургут ихтисослаштирилган давлат ўрмон хўжаликларида": "AP",
    "Ёввойи ҳайвонларни такрор кўпайтириш ва урчитиш питомнигини ташкил этиш учун": "AQ",
    "бошқа фаолият турлари": "AR",
})

LAND_COLUMNS = _columns({
    "Чўл": "AS",
    "Тоғ": "AT",
    "Водий": "AU",
    "Тўқай": "AV",
    "Ўзлаштирилмаган бошқа ерлар": "AW",
})

PRIVILEGE_COLUMNS = _columns({
    # shablon sarlavhalari
    "чўл ўсимликлари (саксовул, черкез, қандим, қорабуроқ) экиш (га) (10 йил)": "AY",
    "ёнғоқ мевали боғлар (писта, бодом, ёнғоқ) ва тез ўсувчи дарахтлар (павловния, терак) плантациялари (5 йил) (га)": "AZ",
    "қишлоқ хўжалиги озиқ-овқат экинлари (дуккакли, мойли, бошоқли экинлар, сабзавот, полиз) (1 йил) (га)": "BA",
    "доривор ўсимликлар плантациялари (3 йил) (га)": "BB",
    # ma'lumotnomadagi (billing) nomlari
    "Ёнғоқ мевали боғлар (писта, бодм, ёнғоқ) тез ўсувчи дарахтлар (павлония ва терак) эгалаган майдонлар (га) (5 йил)": "AZ",
    "Доривор ўсимликлар етиштириш (га) (3 йил)": "BB",
    "Қ/Х экинлари экиш (га) (1 йил)": "BA",
})

COL = {name: column_index_from_string(letter) for name, letter in {
    "row_no": "A",
    "xojalik_inn": "B",
    "xojalik_name": "C",
    "cadastral_number": "D",
    "auction_date": "E",
    "contract_date": "F",
    "expired_date": "G",
    "days_remaining": "H",
    "contract_number": "I",
    "entity_type": "J",
    "stir_or_pinfl": "K",
    "full_name": "L",
    "address": "M",
    "passport": "N",
    "phone": "O",
    "bolim": "P",
    "region_name": "Q",
    "region_code": "R",
    "district_name": "S",
    "district_code": "T",
    "circle": "U",
    "section": "V",
    "counter": "W",
    "massive": "X",
    "rent_area": "Y",
    "privilege_total": "AX",
    "rent_amount": "BC",
    "penalty_amount": "BD",
    "paid_amount": "BE",
    "debt": "BF",
    "vacancy": "BG",
    "own_investment": "BI",
    "product_export": "BJ",
    "product_amount": "BK",
    "export_amount": "BL",
    "canceled_date": "BM",
    "cancel_reason": "BN",
    "latitude": "BO",
    "longitude": "BP",
}.items()}

YEARLY_START_COL = {
    "contract_amount": column_index_from_string("BQ"),
    "calculated_amount": column_index_from_string("BX"),
    "paid_amount": column_index_from_string("CE"),
    "penalty_amount": column_index_from_string("CL"),
}

SUMMABLE_COLUMNS = (
    {COL["rent_area"], COL["privilege_total"], COL["rent_amount"], COL["penalty_amount"], COL["paid_amount"], COL["debt"]}
    | set(ACTIVITY_COLUMNS.values()) | set(LAND_COLUMNS.values()) | set(PRIVILEGE_COLUMNS.values())
    | {start + i for start in YEARLY_START_COL.values() for i in range(len(YEARS))}
)


def export_queryset(queryset: QuerySet) -> QuerySet:
    return queryset.select_related("xojalik", "bolim", "renter", "region", "district").prefetch_related(
        Prefetch("activities", queryset=ContractActivity.objects.select_related("type")),
        Prefetch("lands", queryset=ContractLand.objects.select_related("type")),
        Prefetch("privileges", queryset=ContractPrivilege.objects.select_related("type")),
        Prefetch("yearly_amounts", queryset=ContractYearlyAmount.objects.all()),
        "actions__products",
    )


def _write_typed(ws, row: int, items, columns: dict):
    for item in items:
        col = columns.get(_norm(item.type.name))
        if col:
            ws.cell(row=row, column=col, value=item.area)


def build_contracts_workbook(queryset: QuerySet) -> bytes:
    wb = load_workbook(TEMPLATE_PATH)
    ws = wb[SHEET_NAME]
    template_height = ws.row_dimensions[FIRST_DATA_ROW].height

    row = FIRST_DATA_ROW
    for index, contract in enumerate(export_queryset(queryset).order_by("xojalik__name", "contract_date"), start=1):
        renter = contract.renter
        _write_typed(ws, row, contract.activities.all(), ACTIVITY_COLUMNS)
        _write_typed(ws, row, contract.lands.all(), LAND_COLUMNS)
        privileges = list(contract.privileges.all())
        _write_typed(ws, row, privileges, PRIVILEGE_COLUMNS)
        if privileges:
            ws.cell(row=row, column=COL["privilege_total"], value=sum(p.area for p in privileges))

        values = {
            "row_no": index,
            "xojalik_inn": contract.xojalik.inn,
            "xojalik_name": contract.xojalik.name,
            "cadastral_number": contract.cadastral_number,
            "auction_date": contract.auction_date,
            "contract_date": contract.contract_date,
            "expired_date": contract.expired_date,
            "days_remaining": contract.days_remaining,
            "contract_number": contract.contract_number,
            "entity_type": "Юридик" if renter.is_legal_entity else "Жисмоний",
            "stir_or_pinfl": renter.stir if renter.is_legal_entity else (renter.pinfl or renter.stir),
            "full_name": renter.full_name,
            "address": renter.per_adr,
            "passport": renter.passport,
            "phone": renter.phone,
            "bolim": contract.bolim.name if contract.bolim else None,
            "region_name": contract.region.name if contract.region else "",
            "region_code": contract.region.code if contract.region else "",
            "district_name": contract.district.name if contract.district else "",
            "district_code": contract.district.code if contract.district else "",
            "circle": contract.circle,
            "section": contract.section,
            "counter": contract.counter,
            "massive": contract.massive,
            "rent_area": contract.rent_area,
            "rent_amount": contract.rent_amount,
            "penalty_amount": contract.penalty_amount,
            "paid_amount": contract.paid_amount,
            "debt": contract.debt,
            "canceled_date": contract.canceled_date,
            "cancel_reason": contract.cancel_reason,
            "latitude": float(contract.latitude) if contract.latitude is not None else None,
            "longitude": float(contract.longitude) if contract.longitude is not None else None,
        }
        action = next(iter(contract.actions.all()), None)
        if action:
            values["vacancy"] = action.vacancy
            values["own_investment"] = action.own_investment
            product = next(iter(action.products.all()), None)
            if product:
                values["product_export"] = product.product_export
                values["product_amount"] = product.product_amount
                values["export_amount"] = product.export_amount
        for name, value in values.items():
            ws.cell(row=row, column=COL[name], value=value)

        yearly_by_year = {ya.year: ya for ya in contract.yearly_amounts.all()}
        for i, year in enumerate(YEARS):
            yearly = yearly_by_year.get(year)
            if yearly:
                for field, start in YEARLY_START_COL.items():
                    ws.cell(row=row, column=start + i, value=getattr(yearly, field))

        row += 1
        ws.row_dimensions[row].height = template_height

    last_data_row = row - 1
    if last_data_row >= FIRST_DATA_ROW:
        ws.cell(row=row, column=COL["xojalik_name"], value="ЖАМИ")
        for col in sorted(SUMMABLE_COLUMNS):
            letter = get_column_letter(col)
            ws.cell(row=row, column=col, value=f"=SUM({letter}{FIRST_DATA_ROW}:{letter}{last_data_row})")

    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
