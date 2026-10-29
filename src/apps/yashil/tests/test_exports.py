"""Hisobot eksporti (TZ §14.2, §17.2–17.3; Q21, Q22): XLSX/CSV/PDF, sinxron/fon, formula xavfsizligi, saqlash muddati."""
import csv
import io
from datetime import timedelta
from unittest.mock import patch

from django.core.files.base import ContentFile
from django.test import override_settings
from django.utils import timezone
from openpyxl import load_workbook

from apps.yashil.models import ExportJob, Order, YashilNotification
from apps.yashil.services import exports
from apps.yashil.tasks import export_cleanup
from apps.yashil.tests._common import BASE, auth
from apps.yashil.tests.test_registry import RegistryBase

EXPORTS = f"{BASE}/reports/exports/"


class ExportBase(RegistryBase):
    def setUp(self):
        super().setUp()
        # Formula sifatida bajarilishi mumkin bo'lgan matn (TZ 17.2)
        Order.objects.filter(pk=self.o1.pk).update(title="=HYPERLINK(\"http://x\")")

    def post(self, user, cab, payload, expect=200):
        self.client.credentials(); auth(self.client, user, cab)
        res = self.client.post(EXPORTS, payload, format="json")
        self.assertEqual(res.status_code, expect, getattr(res, "data", res.content))
        return res.data

    def job_bytes(self, data) -> bytes:
        job = ExportJob.objects.get(pk=data["id"])
        with job.file.open("rb") as f:
            return f.read()


class ExportSyncTest(ExportBase):
    def test_xlsx_has_header_block_rows_and_text_safe_formula(self):
        data = self.post(self.o_all, self.o_all_cab, {"kind": "registry", "format": "xlsx", "filters": {"route": "forest"}})
        self.assertEqual(data["status"], "done")
        self.assertEqual(data["row_count"], 3)
        self.assertTrue(data["url"].startswith("/media/documents/yashil/exports/"))
        self.assertIn("?t=", data["url"])
        wb = load_workbook(io.BytesIO(self.job_bytes(data)))
        ws = wb.active
        head = [ws.cell(row=r, column=1).value for r in range(1, 6)]
        self.assertIn("Яшил маҳалла", head[0])
        self.assertTrue(any("Ўрмон хўжалиги" in str(h) for h in head))
        header_row = next(r for r in range(1, 10) if ws.cell(row=r, column=1).value == "Т/р")
        headers = [c.value for c in ws[header_row]]
        self.assertIn("Маҳалла номи", headers)
        self.assertIn("Маҳсулот ёки хизмат", headers)
        title_col = headers.index("Буюртма номи ва рақами") + 1
        cells = [ws.cell(row=r, column=title_col) for r in range(header_row + 1, ws.max_row + 1)]
        formula_like = [c for c in cells if "=HYPERLINK" in str(c.value)]
        self.assertEqual(len(formula_like), 2)  # o1 ning ikkala qatori
        self.assertTrue(all(c.data_type == "s" for c in formula_like))
        self.assertEqual(ws.max_row, header_row + 3)

    def test_csv_utf8_bom_and_formula_guard(self):
        data = self.post(self.o_all, self.o_all_cab, {"kind": "registry", "format": "csv", "filters": {}})
        raw = self.job_bytes(data)
        self.assertTrue(raw.startswith(b"\xef\xbb\xbf"))
        rows = list(csv.reader(io.StringIO(raw.decode("utf-8-sig"))))
        header_i = next(i for i, r in enumerate(rows) if r and r[0] == "Т/р")
        body = rows[header_i + 1:]
        self.assertEqual(len(body), 4)
        title_col = rows[header_i].index("Буюртма номи ва рақами")
        bad = [r[title_col] for r in body if "HYPERLINK" in r[title_col]]
        self.assertTrue(bad and all(v.startswith("'") for v in bad))

    def test_pdf_is_generated(self):
        data = self.post(self.o_all, self.o_all_cab, {"kind": "registry", "format": "pdf", "filters": {"state": ["execution"]}})
        raw = self.job_bytes(data)
        self.assertTrue(raw.startswith(b"%PDF"))
        self.assertEqual(data["row_count"], 2)

    def test_acts_kind(self):
        data = self.post(self.o_all, self.o_all_cab, {"kind": "acts", "format": "xlsx", "filters": {}})
        self.assertEqual(data["row_count"], 2)
        ws = load_workbook(io.BytesIO(self.job_bytes(data))).active
        header_row = next(r for r in range(1, 10) if ws.cell(row=r, column=1).value == "Т/р")
        self.assertIn("Далолатнома рақами", [c.value for c in ws[header_row]])

    def test_scope_mahalla_and_observer_can_export(self):
        data = self.post(self.m_user, self.m_cab, {"kind": "registry", "format": "csv", "filters": {}})
        self.assertEqual(data["row_count"], 3)
        data = self.post(self.o_user, self.o_cab, {"kind": "registry", "format": "csv", "filters": {}})
        self.assertEqual(data["row_count"], 3)

    def test_invalid_format_and_filters(self):
        res = self.post(self.o_all, self.o_all_cab, {"kind": "registry", "format": "doc", "filters": {}}, expect=400)
        self.assertEqual(res["error_code"], "invalid_format")
        res = self.post(self.o_all, self.o_all_cab, {"kind": "registry", "format": "csv", "filters": {"period_basis": "zzz"}}, expect=400)
        self.assertEqual(res["field"], "period_basis")

    @override_settings(YASHIL_EXPORT_PDF_LIMIT=1)
    def test_pdf_over_limit_is_422(self):
        res = self.post(self.o_all, self.o_all_cab, {"kind": "registry", "format": "pdf", "filters": {}}, expect=422)
        self.assertEqual(res["error_code"], "export_too_large")


class ExportAsyncTest(ExportBase):
    @override_settings(YASHIL_EXPORT_SYNC_LIMIT=1)
    def test_large_export_is_queued_then_done_with_notification(self):
        with patch("apps.yashil.tasks.export_job.delay") as delay, self.captureOnCommitCallbacks(execute=True):
            data = self.post(self.o_all, self.o_all_cab, {"kind": "registry", "format": "xlsx", "filters": {}}, expect=202)
        self.assertEqual(data["status"], "queued")
        self.assertIsNone(data["url"])
        delay.assert_called_once_with(data["id"])
        job = ExportJob.objects.get(pk=data["id"])
        exports.run_export(job)
        job.refresh_from_db()
        self.assertEqual(job.status, "done")
        self.assertEqual(job.row_count, 4)
        n = YashilNotification.objects.filter(recipient_user=self.o_all, kind="export_ready")
        self.assertEqual(n.count(), 1)
        self.client.credentials(); auth(self.client, self.o_all, self.o_all_cab)
        res = self.client.get(f"{EXPORTS}{job.pk}/")
        self.assertEqual(res.status_code, 200)
        self.assertEqual(res.data["status"], "done")
        self.assertIn("?t=", res.data["url"])

    @override_settings(YASHIL_EXPORT_SYNC_LIMIT=1, YASHIL_EXPORT_MAX_ACTIVE=2)
    def test_too_many_active_jobs_is_429(self):
        with patch("apps.yashil.tasks.export_job.delay"):
            for _ in range(2):
                self.post(self.o_all, self.o_all_cab, {"kind": "registry", "format": "csv", "filters": {}}, expect=202)
            res = self.post(self.o_all, self.o_all_cab, {"kind": "registry", "format": "csv", "filters": {}}, expect=429)
        self.assertEqual(res["error_code"], "too_many_exports")

    def test_failed_job_records_error_and_notifies(self):
        job = ExportJob.objects.create(requested_by=self.o_all, membership=self.o_all_cab, kind="registry", format="csv", filters={})
        with patch("apps.yashil.services.exports._write", side_effect=RuntimeError("disk")):
            exports.run_export(job)
        job.refresh_from_db()
        self.assertEqual(job.status, "failed")
        self.assertIn("disk", job.error)
        self.assertEqual(YashilNotification.objects.filter(recipient_user=self.o_all, kind="export_failed").count(), 1)

    def test_other_users_job_is_404_and_list_is_own(self):
        data = self.post(self.o_all, self.o_all_cab, {"kind": "registry", "format": "csv", "filters": {}})
        self.client.credentials(); auth(self.client, self.m_user, self.m_cab)
        self.assertEqual(self.client.get(f"{EXPORTS}{data['id']}/").status_code, 404)
        self.assertEqual(self.client.get(EXPORTS).data["results"], [])
        self.client.credentials(); auth(self.client, self.o_all, self.o_all_cab)
        self.assertEqual([j["id"] for j in self.client.get(EXPORTS).data["results"]], [data["id"]])


class ExportCleanupTest(ExportBase):
    def test_cleanup_removes_old_files_only(self):
        fresh = self.post(self.o_all, self.o_all_cab, {"kind": "registry", "format": "csv", "filters": {}})
        old = ExportJob.objects.create(requested_by=self.o_all, membership=self.o_all_cab, kind="registry", format="csv", filters={},
                                       status="done", row_count=1)
        old.file.save("old.csv", ContentFile(b"x"))
        ExportJob.objects.filter(pk=old.pk).update(finished_at=timezone.now() - timedelta(days=8))
        removed = export_cleanup()
        self.assertEqual(removed, 1)
        old.refresh_from_db()
        self.assertEqual(old.status, "expired")
        self.assertFalse(old.file)
        self.assertEqual(ExportJob.objects.get(pk=fresh["id"]).status, "done")
