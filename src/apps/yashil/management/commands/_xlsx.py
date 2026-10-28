"""
Excel import yordamchilari: sarlavhalarni (kirill/lotin) tanish, sana/mantiqiy
qiymatlarni o'qish, qator xatolarini yig'ish. Fayl avval `excel_safety` bilan tekshiriladi.
"""
import re
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path

from django.core.management.base import CommandError
from openpyxl import load_workbook

from apps.utils.excel_safety import UnsafeExcelFile, check_xlsx_zip

TRUE_WORDS = {"ҳа", "ha", "xa", "да", "yes", "true", "1", "faol", "фаол"}
FALSE_WORDS = {"йўқ", "yo'q", "yoq", "yo‘q", "нет", "no", "false", "0", "nofaol", "нофаол"}


def norm_header(value) -> str:
    """Sarlavha: kichik harf, apostrof/bo'shliqlarsiz, kirill 'ҳ'→'h' va h.k. solishtirish uchun."""
    s = str(value or "").strip().lower()
    s = re.sub(r"[‘’ʻʼ'`´\s_\-]+", "", s)
    table = str.maketrans({
        "а": "a", "б": "b", "в": "v", "г": "g", "д": "d", "е": "e", "ё": "yo", "ж": "j", "з": "z", "и": "i",
        "й": "y", "к": "k", "л": "l", "м": "m", "н": "n", "о": "o", "п": "p", "р": "r", "с": "s", "т": "t",
        "у": "u", "ф": "f", "х": "x", "ц": "ts", "ч": "ch", "ш": "sh", "щ": "sh", "ъ": "", "ь": "", "э": "e",
        "ю": "yu", "я": "ya", "ў": "o", "қ": "q", "ғ": "g", "ҳ": "h",
    })
    return s.translate(table)


def parse_date(value) -> date | None:
    if value in (None, ""):
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    s = str(value).strip()
    for fmt in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            continue
    raise ValueError(f"Sana tushunarsiz: {s!r} (dd.mm.yyyy yoki yyyy-mm-dd)")


def parse_bool(value, default=True) -> bool:
    if value in (None, ""):
        return default
    s = str(value).strip().lower()
    if s in TRUE_WORDS:
        return True
    if s in FALSE_WORDS:
        return False
    raise ValueError(f"Mantiqiy qiymat tushunarsiz: {value!r}")


def text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


@dataclass
class Sheet:
    headers: list
    rows: list  # (row_number, dict[normalized_header -> value])
    errors: list = field(default_factory=list)

    def col(self, row: dict, *aliases, required=False, row_no=None):
        for alias in aliases:
            key = norm_header(alias)
            if key in row and row[key] not in (None, ""):
                return row[key]
        if required:
            raise ValueError(f"'{aliases[0]}' ustuni bo'sh")
        return None


def read_sheet(path: str, required_headers: dict) -> Sheet:
    """
    `required_headers`: {"maqsad": ["манба коди", "manba kodi", ...]} — har maqsad uchun
    kamida bitta alias faylda bo'lishi shart.
    """
    file_path = Path(path)
    if not file_path.exists():
        raise CommandError(f"Fayl topilmadi: {path}")
    with file_path.open("rb") as fh:
        try:
            check_xlsx_zip(fh)
        except UnsafeExcelFile as exc:
            raise CommandError(f"Excel fayl xavfsizlik tekshiruvidan o'tmadi: {exc}")
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    iterator = ws.iter_rows(values_only=True)
    try:
        header_row = next(iterator)
    except StopIteration:
        raise CommandError("Fayl bo'sh")
    headers = [norm_header(h) for h in header_row]
    missing = [aliases[0] for aliases in required_headers.values() if not any(norm_header(a) in headers for a in aliases)]
    if missing:
        raise CommandError("Majburiy ustunlar topilmadi: " + ", ".join(missing))
    rows = []
    for idx, values in enumerate(iterator, start=2):
        if values is None or all(v in (None, "") for v in values):
            continue
        rows.append((idx, {h: v for h, v in zip(headers, values) if h}))
    wb.close()
    return Sheet(headers=headers, rows=rows)


class ImportReport:
    def __init__(self, stdout, dry_run: bool):
        self.stdout = stdout
        self.dry_run = dry_run
        self.created = self.updated = self.unchanged = 0
        self.errors = []

    def error(self, row_no, message):
        self.errors.append(f"{row_no}-qator: {message}")

    def print(self):
        prefix = "[DRY-RUN] " if self.dry_run else ""
        self.stdout.write(f"{prefix}yangi: {self.created}, yangilandi: {self.updated}, o'zgarishsiz: {self.unchanged}, xato: {len(self.errors)}")
        for line in self.errors:
            self.stdout.write("  " + line)
