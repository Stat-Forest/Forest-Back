"""
Qabul sinovlari bayonnomasi (TZ §18: rol, kiritilgan ma'lumot, kutilgan va amaldagi natija).

`run_acceptance_tests` — katalogdagi avtomatik testlarni alohida test bazasida ishga tushiradi
va har test natijasini qaytaradi. `build_protocol_docx` — DOCX bayonnoma: avtomatik natija
to'ldiriladi, «Амалдаги натижа (қўлда)» ustuni test serverda sinov o'tkazganda to'ldiriladi.
"""
import time
import unittest
from datetime import datetime

from apps.yashil.acceptance import CASES, all_test_labels

OK, FAIL, ERROR, SKIP, MISSING = "ok", "fail", "error", "skip", "not_run"
RESULT_LABELS = {OK: "ўтди", FAIL: "ўтмади", ERROR: "хато", SKIP: "ўтказиб юборилди", MISSING: "ишга туширилмади"}


class _RecordingResult(unittest.TextTestResult):
    records: dict = {}

    def startTest(self, test):
        self._started = time.monotonic()
        super().startTest(test)

    def _record(self, test, status, detail=""):
        label = test.id()
        type(self).records[label] = {"status": status, "seconds": round(time.monotonic() - getattr(self, "_started", time.monotonic()), 2),
                                     "detail": detail[-2000:]}

    def addSuccess(self, test):
        super().addSuccess(test)
        self._record(test, OK)

    def addFailure(self, test, err):
        super().addFailure(test, err)
        self._record(test, FAIL, self._exc_info_to_string(err, test))

    def addError(self, test, err):
        super().addError(test, err)
        self._record(test, ERROR, self._exc_info_to_string(err, test))

    def addSkip(self, test, reason):
        super().addSkip(test, reason)
        self._record(test, SKIP, reason)


def run_acceptance_tests(verbosity: int = 1) -> dict[str, dict]:
    from django.test.runner import DiscoverRunner

    class Runner(DiscoverRunner):
        def get_resultclass(self):
            return _RecordingResult

    _RecordingResult.records = {}
    Runner(verbosity=verbosity, interactive=False).run_tests(all_test_labels())
    return dict(_RecordingResult.records)


def case_summary(case, results: dict[str, dict]) -> dict:
    labels = [f"apps.yashil.tests.{t}" for t in case.tests]
    statuses = [results.get(l, {}).get("status", MISSING) for l in labels]
    if not labels:
        auto = "—"
        status = "manual"
    else:
        passed = sum(1 for s in statuses if s == OK)
        auto = f"{passed}/{len(labels)} {RESULT_LABELS[OK]}" if passed == len(labels) else f"{passed}/{len(labels)} — {RESULT_LABELS[FAIL]}"
        status = OK if passed == len(labels) else FAIL
    return {"id": case.id, "auto": auto, "status": status, "tests": list(zip(labels, statuses))}


def build_protocol_docx(results: dict[str, dict] | None, meta: dict):
    from docx import Document
    from docx.enum.section import WD_ORIENT
    from docx.shared import Cm, Pt

    doc = Document()
    section = doc.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width, section.page_height = section.page_height, section.page_width
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(section, side, Cm(1.5))
    style = doc.styles["Normal"]
    style.font.name, style.font.size = "Times New Roman", Pt(10)

    doc.add_heading("«Яшил маҳалла» модули — қабул синовлари баённомаси (Q01–Q36)", level=1)
    p = doc.add_paragraph()
    p.add_run("Асос: ").bold = True
    p.add_run("техник топшириқ 8.0, §18–19, §23.1. Ҳар синовда бажарган роль, киритилган маълумот, кутилган ва амалдаги натижа қайд этилади.")
    for key, label in (("date", "Сана"), ("environment", "Муҳит"), ("version", "Дастур версияси (commit)"), ("note", "Изоҳ")):
        if meta.get(key):
            q = doc.add_paragraph()
            q.add_run(f"{label}: ").bold = True
            q.add_run(str(meta[key]))

    summaries = {c.id: case_summary(c, results or {}) for c in CASES}
    if results is not None:
        auto_cases = [s for s in summaries.values() if s["status"] != "manual"]
        passed = sum(1 for s in auto_cases if s["status"] == OK)
        q = doc.add_paragraph()
        q.add_run("Автоматик синовлар: ").bold = True
        q.add_run(f"{len(auto_cases)} та ҳолатдан {passed} тасида ҳамма тестлар ўтди; ишга туширилган тестлар сони {len(results)}.")

    headers = ("ID", "Синов", "Роль", "Киритилган маълумот", "Кутилган натижа", "Автоматик синов", "Амалдаги натижа (қўлда)", "Изоҳ / имзо")
    widths = (Cm(1.2), Cm(4.2), Cm(2.6), Cm(3.0), Cm(4.4), Cm(2.4), Cm(4.4), Cm(2.6))
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.autofit = False
    for col, w in zip(table.columns, widths):
        col.width = w
    for cell, text, w in zip(table.rows[0].cells, headers, widths):
        cell.text = ""
        cell.paragraphs[0].add_run(text).bold = True
        cell.width = w
    for c in CASES:
        s = summaries[c.id]
        auto = s["auto"] if results is not None else ("—" if not c.tests else f"{len(c.tests)} та тест")
        manual_hint = c.manual + (f" Восита: {c.tools}." if c.tools else "")
        row = table.add_row().cells
        for cell, text, w in zip(row, (c.id, c.title, c.role, c.data, c.expected, auto, f"Қадам: {manual_hint}\nНатижа: ", ""), widths):
            cell.text = text
            cell.width = w

    doc.add_heading("Илова: автоматик тестлар", level=2)
    for c in CASES:
        if not c.tests:
            continue
        doc.add_paragraph(f"{c.id}", style="List Bullet")
        for label, status in summaries[c.id]["tests"]:
            shown = RESULT_LABELS[status] if results is not None else "—"
            doc.add_paragraph(f"{label.removeprefix('apps.yashil.tests.')} — {shown}", style="List Bullet 2")

    doc.add_heading("Тасдиқлаш", level=2)
    for line in ("Буюртмачи ташкилот: ______________________________", "Масъул шахс ва лавозим: ___________________________",
                 "Ишлаб чиқувчи: ___________________________________", "Сана: ____________________"):
        doc.add_paragraph(line)
    return doc


def default_meta(environment: str = "", version: str = "", note: str = "") -> dict:
    return {"date": datetime.now().strftime("%d.%m.%Y %H:%M"), "environment": environment, "version": version, "note": note}
