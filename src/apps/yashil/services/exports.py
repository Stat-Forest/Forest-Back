"""
Hisobot eksporti (TZ §14.2, §17.2–17.3; javob З1).

    create_job(user, cabinet, is_admin, kind, fmt, filters) -> ExportJob
        qatorlar soni ≤ YASHIL_EXPORT_SYNC_LIMIT — shu yerda tayyorlanadi (status done),
        aks holda queued + Celery `yashil.export_job`.
    run_export(job)                                        -> ExportJob
    cleanup(retention_days)                                -> int

Fayl: XLSX (openpyxl, matn xavfsiz — `excel_safe`), CSV (UTF-8 BOM; `=+-@` bilan
boshlangan matn oldiga `'`), PDF (reportlab landshaft, sarlavha takrori).
Faylning boshida sarlavha bloki: modul, hisobot turi, filtrlar, kabinet, vaqt.
Ko'lam — `registry` / `scope` servislari (rol va hudud serverda).
"""
import csv
import io
import logging
from datetime import timedelta

from django.core.files.base import ContentFile
from django.db import transaction
from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from apps.utils.excel_safe import append_safe
from apps.yashil import conf
from apps.yashil.models import STATE_LABELS_CYR, Act, ExportJob, YashilNotification
from apps.yashil.services import registry
from apps.yashil.services.scope import acts_for_cabinet
from apps.yashil.services.units import UNIT_LABELS

log = logging.getLogger(__name__)
K = YashilNotification.Kind

# (kalit, sarlavha) — TZ 14.1 majburiy ustunlar; keyin tanlanadigan qo'shimchalar (javob З2)
REGISTRY_COLUMNS = [
    ("row_no", "Т/р"),
    ("region", "Вилоят"),
    ("district", "Туман"),
    ("mahalla", "Маҳалла номи"),
    ("supplier", "Етказиб берувчи"),
    ("route_cyr", "Йўналиш"),
    ("number", "Буюртма рақами"),
    ("title", "Буюртма номи ва рақами"),
    ("item", "Маҳсулот ёки хизмат"),
    ("qty", "Миқдор"),
    ("unit_cyr", "Бирлик"),
    ("accepted_qty", "Қабул қилинган"),
    ("remaining_qty", "Қолган"),
    ("state_cyr", "Буюртма ҳолати"),
    ("responsible_cyr", "Масъул босқич"),
    ("first_submitted_on", "Буюртма берилган сана"),
    ("last_submitted_at", "Охирги қайта юбориш"),
    ("review_due", "Кўриб чиқиш муддати"),
    ("delivery_due", "Етказиш муддати"),
    ("late_workdays", "Кечиккан иш кунлари"),
    ("overdue_stage_cyr", "Кечикиш босқичи"),
]
REGISTRY_OPTIONAL = [
    ("executor_name", "Ижрочи"),
    ("act_numbers", "Далолатнома рақамлари"),
    ("last_delivery_date", "Ҳақиқий етказиш санаси"),
    ("accepted_on", "Ижрога қабул санаси"),
    ("address", "Манзил"),
    ("lat", "Кенглик"),
    ("lon", "Узунлик"),
    ("area_ha", "Майдон, га"),
]
ACT_COLUMNS = [
    ("row_no", "Т/р"),
    ("number", "Далолатнома рақами"),
    ("version_no", "Версия"),
    ("status_cyr", "Ҳолат"),
    ("order_number", "Буюртма рақами"),
    ("order_state_cyr", "Буюртма ҳолати"),
    ("mahalla", "Маҳалла"),
    ("supplier", "Етказиб берувчи"),
    ("route_cyr", "Йўналиш"),
    ("delivery_date", "Ҳақиқий етказиш санаси"),
    ("submitted_at", "Етказиб берувчи имзоси"),
    ("approved_at", "Маҳалла тасдиғи"),
    ("is_partial", "Қисман"),
    ("responsible", "Масъул шахс"),
    ("lines", "Позициялар"),
]
ACT_STATUS_CYR = {
    Act.Status.DRAFT: "Қоралама", Act.Status.REVIEW: "Маҳалла тасдиғида", Act.Status.APPROVED: "Қабул қилинган",
    Act.Status.RETURNED: "Тузатиш учун қайтарилган", Act.Status.CANCELLED: "Бекор қилинган",
}
KIND_TITLES = {ExportJob.Kind.REGISTRY: "Буюртмалар реестри", ExportJob.Kind.ACTS: "Далолатномалар реестри"}
CONTENT_TYPES = {
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "csv": "text/csv; charset=utf-8",
    "pdf": "application/pdf",
}


class ExportError(Exception):
    def __init__(self, message, *, code="business_rule", status=422, field=None):
        super().__init__(message)
        self.message, self.code, self.status, self.field = message, code, status, field


# ----------------------------------------------------------------------------- qatorlar
def _registry_columns(filters: dict):
    extra = set(filters.get("columns") or [])
    return REGISTRY_COLUMNS + [c for c in REGISTRY_OPTIONAL if c[0] in extra]


def _registry_rows(job: ExportJob, f: dict):
    qs = registry.registry_queryset(job.membership, job.membership is None, f)
    chunk = conf.get("YASHIL_EXPORT_CHUNK")
    start = 1
    buffer = []
    for line in qs.iterator(chunk_size=chunk):
        buffer.append(line)
        if len(buffer) >= chunk:
            yield from registry.row_dicts(buffer, start_no=start)
            start += len(buffer)
            buffer = []
    if buffer:
        yield from registry.row_dicts(buffer, start_no=start)


def act_queryset(cabinet, is_admin: bool, f: dict):
    """Aktlar reestri (javob З2 — alohida reestr); filtrlar reestr lug'ati bilan bir xil nomda."""
    qs = acts_for_cabinet(cabinet, is_admin).select_related("order__mahalla__district__region", "order__supplier_org", "current_version")
    if f.get("date_from"):
        qs = qs.filter(current_version__delivery_date__gte=f["date_from"])
    if f.get("date_to"):
        qs = qs.filter(current_version__delivery_date__lte=f["date_to"])
    if f.get("region"):
        qs = qs.filter(order__mahalla__region_id=f["region"])
    if f.get("district"):
        qs = qs.filter(order__mahalla__district_id=f["district"])
    if f.get("mahalla"):
        qs = qs.filter(order__mahalla_id=f["mahalla"])
    if f.get("supplier_org"):
        qs = qs.filter(order__supplier_org_id=f["supplier_org"])
    if f.get("route"):
        qs = qs.filter(order__route=f["route"])
    if f.get("status"):
        qs = qs.filter(status__in=f["status"])
    if f.get("q"):
        q = f["q"]
        qs = qs.filter(number__icontains=q) | qs.filter(order__title__icontains=q) | qs.filter(order__mahalla__name__icontains=q)
    return qs.order_by("-updated_at", "number")


def act_rows(acts, start_no: int = 1) -> list[dict]:
    acts = list(acts)
    rows = []
    for n, a in enumerate(acts, start=start_no):
        o, v = a.order, a.current_version
        lines = ""
        if v is not None:
            parts = []
            for l in v.lines.select_related("order_line__catalog_item"):
                item = l.order_line.catalog_item
                parts.append(f"{item.name_cyr} — {format(l.qty.normalize(), 'f')} {UNIT_LABELS.get(item.unit, (item.unit,))[0]}")
            lines = "; ".join(parts)
        rows.append({
            "row_no": n, "id": str(a.id), "number": a.number, "version_no": a.version_no, "status": a.status,
            "status_cyr": ACT_STATUS_CYR.get(a.status, a.status), "order_id": str(o.id), "order_number": o.number,
            "order_state_cyr": STATE_LABELS_CYR.get(o.state, o.state), "mahalla": o.mahalla.name, "supplier": o.supplier_org.name,
            "route_cyr": registry.ROUTE_CYR.get(o.route, o.route),
            "delivery_date": v.delivery_date.isoformat() if v else None,
            "submitted_at": timezone.localtime(v.submitted_at).strftime("%d.%m.%Y %H:%M") if v and v.submitted_at else None,
            "approved_at": timezone.localtime(a.approved_at).strftime("%d.%m.%Y %H:%M") if a.approved_at else None,
            "is_partial": "ҳа" if a.is_partial else "йўқ", "responsible": v.responsible if v else "", "lines": lines,
        })
    return rows


def _act_rows(job: ExportJob, f: dict):
    qs = act_queryset(job.membership, job.membership is None, f)
    chunk = conf.get("YASHIL_EXPORT_CHUNK")
    start, buffer = 1, []
    for act in qs.iterator(chunk_size=chunk):
        buffer.append(act)
        if len(buffer) >= chunk:
            yield from act_rows(buffer, start_no=start)
            start += len(buffer)
            buffer = []
    if buffer:
        yield from act_rows(buffer, start_no=start)


def parse_job_filters(kind: str, raw: dict) -> dict:
    """Eksport `filters` JSON'i — reestr parametrlari bilan bir xil tekshiruv."""
    raw = dict(raw or {})
    flat = {}
    for key, value in raw.items():
        if isinstance(value, (list, tuple)):
            flat[key] = ",".join(str(v) for v in value)
        elif isinstance(value, bool):
            flat[key] = "1" if value else ""
        elif value is not None:
            flat[key] = str(value)
    if kind == ExportJob.Kind.ACTS:
        flat.pop("state", None)
    f = registry.parse_filters(flat)
    if kind == ExportJob.Kind.ACTS:
        statuses = registry._csv(flat.get("status"))
        bad = [s for s in statuses if s not in Act.Status.values]
        if bad:
            raise registry.RegistryFilterError(f"Номаълум акт ҳолати: {bad[0]}", "status")
        f["status"] = statuses
    cols = raw.get("columns") or []
    allowed = {c[0] for c in REGISTRY_OPTIONAL}
    f["columns"] = [c for c in cols if c in allowed]
    return f


def _serializable(f: dict) -> dict:
    out = {}
    for k, v in f.items():
        out[k] = v.isoformat() if hasattr(v, "isoformat") else v
    return out


def count_rows(job: ExportJob, f: dict) -> int:
    if job.kind == ExportJob.Kind.ACTS:
        return act_queryset(job.membership, job.membership is None, f).count()
    return registry.apply_filters(registry.base_queryset(job.membership, job.membership is None), f).count()


# ----------------------------------------------------------------------------- yozuvchilar
def _header_lines(job: ExportJob, f: dict) -> list[str]:
    cab = job.membership
    who = cab.cabinet_name if cab is not None else "Модул администратори"
    when = timezone.localtime(timezone.now()).strftime("%d.%m.%Y %H:%M")
    return [
        f"STAT-URMON · Яшил маҳалла · {KIND_TITLES.get(job.kind, job.kind)}",
        f"Йўналиш: {registry.ROUTE_CYR.get(f.get('route'), 'Умумий (иккала йўналиш)')}",
        f"Фильтрлар: {registry.filters_label(f)}",
        f"Кабинет: {who} · Шакллантирилди: {when} (Asia/Tashkent) · Қатор: {job.row_count}",
    ]


def _cell(value):
    if value is None:
        return ""
    return value


def _write_xlsx(job, f, columns, rows) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = KIND_TITLES.get(job.kind, "Ҳисобот")[:31]
    for line in _header_lines(job, f):
        append_safe(ws, [line])
        ws.cell(row=ws.max_row, column=1).font = Font(bold=ws.max_row == 1)
    ws.append([])
    append_safe(ws, [h for _, h in columns])
    header_row = ws.max_row
    fill = PatternFill("solid", fgColor="E8EEF2")
    for c in ws[header_row]:
        c.font = Font(bold=True)
        c.fill = fill
        c.alignment = Alignment(wrap_text=True, vertical="top")
    for r in rows:
        append_safe(ws, [_cell(r.get(k)) for k, _ in columns])
    for i, (key, _) in enumerate(columns, start=1):
        ws.column_dimensions[get_column_letter(i)].width = 6 if key == "row_no" else 22
    ws.freeze_panes = ws.cell(row=header_row + 1, column=1)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _csv_safe(value) -> str:
    s = "" if value is None else str(value)
    return "'" + s if s[:1] in ("=", "+", "-", "@") else s


def _write_csv(job, f, columns, rows) -> bytes:
    buf = io.StringIO()
    w = csv.writer(buf, lineterminator="\n")
    for line in _header_lines(job, f):
        w.writerow([_csv_safe(line)])
    w.writerow([])
    w.writerow([h for _, h in columns])
    for r in rows:
        w.writerow([_csv_safe(r.get(k)) for k, _ in columns])
    return b"\xef\xbb\xbf" + buf.getvalue().encode("utf-8")


def _write_pdf(job, f, columns, rows) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.units import mm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    from apps.yashil.services import pdf as pdfs

    pdfs._register_fonts()
    st = pdfs._styles()
    buf = io.BytesIO()
    page = landscape(A4)
    doc = SimpleDocTemplate(buf, pagesize=page, leftMargin=10 * mm, rightMargin=10 * mm, topMargin=12 * mm, bottomMargin=12 * mm,
                            title=KIND_TITLES.get(job.kind, "Ҳисобот"), author="STAT-URMON")
    story = []
    for i, line in enumerate(_header_lines(job, f)):
        story.append(Paragraph(pdfs._esc(line), st["h"] if i == 0 else st["small"]))
    story.append(Spacer(1, 4))
    cell = st["cell"]
    cell.fontSize, cell.leading = 7, 8.5
    cellh = st["cellh"]
    cellh.fontSize, cellh.leading = 7, 8.5
    data = [[Paragraph(h, cellh) for _, h in columns]]
    for r in rows:
        data.append([Paragraph(pdfs._esc(r.get(k)), cell) for k, _ in columns])
    width = page[0] - 20 * mm
    first = 9 * mm
    rest = (width - first) / max(1, len(columns) - 1)
    t = Table(data, colWidths=[first] + [rest] * (len(columns) - 1), repeatRows=1)
    t.setStyle(TableStyle([
        ("GRID", (0, 0), (-1, -1), 0.3, colors.HexColor("#999999")),
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#e8eef2")),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
    ]))
    story.append(t)

    def footer(canvas, d):
        canvas.saveState()
        canvas.setFont(pdfs.FONT, 7)
        canvas.setFillColor(colors.HexColor("#666666"))
        canvas.drawString(10 * mm, 6 * mm, "STAT-URMON · Яшил маҳалла")
        canvas.drawRightString(page[0] - 10 * mm, 6 * mm, f"{d.page}-бет")
        canvas.restoreState()

    doc.build(story, onFirstPage=footer, onLaterPages=footer)
    return buf.getvalue()


WRITERS = {"xlsx": _write_xlsx, "csv": _write_csv, "pdf": _write_pdf}


def _write(job: ExportJob, f: dict) -> tuple[bytes, int]:
    if job.kind == ExportJob.Kind.ACTS:
        columns, rows = ACT_COLUMNS, list(_act_rows(job, f))
    else:
        columns, rows = _registry_columns(f), list(_registry_rows(job, f))
    job.row_count = len(rows)
    return WRITERS[job.format](job, f, columns, rows), len(rows)


# ----------------------------------------------------------------------------- vazifa
def run_export(job: ExportJob) -> ExportJob:
    """Faylni tayyorlaydi, saqlaydi, so'rovchiga xabar beradi. Xato — status failed + xabar."""
    ExportJob.objects.filter(pk=job.pk).update(status=ExportJob.Status.RUNNING, started_at=timezone.now())
    try:
        f = parse_job_filters(job.kind, job.filters)
        data, count = _write(job, f)
        job.file.save(f"{job.id}.{job.format}", ContentFile(data), save=False)
        job.row_count = count
        job.status = ExportJob.Status.DONE
        job.error = ""
    except Exception as exc:  # noqa: BLE001 — har qanday xato vazifada qayd etiladi
        log.exception("yashil export %s failed", job.pk)
        job.status = ExportJob.Status.FAILED
        job.error = str(exc)[:2000]
    job.finished_at = timezone.now()
    job.save(update_fields=["file", "row_count", "status", "error", "finished_at"])
    if getattr(job, "_notify", True):
        _notify(job)
    return job


def _notify(job: ExportJob):
    title = KIND_TITLES.get(job.kind, job.kind)
    if job.status == ExportJob.Status.DONE:
        kind, text, body = K.EXPORT_READY, f"Экспорт тайёр: {title} ({job.format.upper()})", f"{job.row_count} қатор. Ҳисоботлар бўлимидан юклаб олинг."
    else:
        kind, text, body = K.EXPORT_FAILED, f"Экспорт бажарилмади: {title}", "Қайта уриниб кўринг ёки администраторга мурожаат қилинг."
    YashilNotification.objects.create(recipient_user=job.requested_by, membership=job.membership, kind=kind, title=text, body=body)


def create_job(user, cabinet, is_admin: bool, *, kind: str, fmt: str, filters: dict) -> ExportJob:
    """So'rov: kichik — darhol tayyor (xabarsiz), katta — navbat (tayyor bo'lganda xabar)."""
    if kind not in ExportJob.Kind.values:
        raise ExportError("Ҳисобот тури: registry / acts", code="invalid_format", status=400, field="kind")
    if fmt not in ExportJob.Format.values:
        raise ExportError("Формат: xlsx / csv / pdf", code="invalid_format", status=400, field="format")
    try:
        f = parse_job_filters(kind, filters)
    except registry.RegistryFilterError as exc:
        raise ExportError(str(exc), code="invalid_format", status=400, field=exc.field)
    active = ExportJob.objects.filter(requested_by=user, status__in=ExportJob.ACTIVE_STATUSES).count()
    if active >= conf.get("YASHIL_EXPORT_MAX_ACTIVE"):
        raise ExportError("Бир вақтда жуда кўп экспорт — олдингилари тайёр бўлишини кутинг", code="too_many_exports", status=429)
    job = ExportJob(requested_by=user, membership=None if (is_admin and cabinet is None) else cabinet, kind=kind, format=fmt,
                    filters=_serializable({**f, "columns": f.get("columns", [])}))
    total = count_rows(job, f)
    if fmt == ExportJob.Format.PDF and total > conf.get("YASHIL_EXPORT_PDF_LIMIT"):
        raise ExportError(f"PDF учун жуда кўп қатор ({total}). Фильтрни қисқартиринг ёки XLSX/CSV танланг.", code="export_too_large")
    job.save()
    if total <= conf.get("YASHIL_EXPORT_SYNC_LIMIT"):
        job._notify = False
        return run_export(job)
    from apps.yashil.tasks import export_job

    transaction.on_commit(lambda: export_job.delay(str(job.pk)))
    return job


def cleanup(retention_days: int | None = None) -> int:
    """Muddati o'tgan fayllarni o'chiradi (yozuv qoladi, status expired)."""
    days = retention_days if retention_days is not None else conf.get("YASHIL_EXPORT_RETENTION_DAYS")
    edge = timezone.now() - timedelta(days=days)
    removed = 0
    for job in ExportJob.objects.filter(status=ExportJob.Status.DONE, finished_at__lt=edge):
        if job.file:
            job.file.delete(save=False)
        job.status = ExportJob.Status.EXPIRED
        job.save(update_fields=["file", "status"])
        removed += 1
    return removed
