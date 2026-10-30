"""
Excel importlari (javoblar A1, A2, A4, Д1, Е1; plans §11.1-9, §11.1-10).

    import_mahallas(source, dry_run, partial, user)
    import_assignments(source, dry_run, partial, user, source_doc)
    import_routes(source, dry_run, partial, user, source_doc)
    import_calendar(source, dry_run, partial, user, source_ref)
    import_opening_stock(source, cabinet, dry_run, partial)

`source` — fayl yo'li yoki fayl obyekti (API yuklamasi). Har qator shu tranzaksiyada
yoziladi (oldingi qatorlar ta'siri — masalan davr kesishuvi — ko'rinadi), keyin:
`dry_run` yoki (`partial=False` va xato bor) bo'lsa hammasi bekor qilinadi.
API `partial=False` ishlatadi: fayl to'liq to'g'ri bo'lmaguncha hech narsa yozilmaydi.
Buyruqlar (`manage.py yashil_import_*`) avvalgidek xatoli qatorlarni tashlab o'tadi.
"""
import re
import uuid
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path

from django.core.exceptions import ValidationError
from django.db import transaction
from openpyxl import load_workbook

from apps.utils.excel_safety import UnsafeExcelFile, check_xlsx_zip

TRUE_WORDS = {"ҳа", "ha", "xa", "да", "yes", "true", "1", "faol", "фаол"}
FALSE_WORDS = {"йўқ", "yo'q", "yoq", "yo‘q", "нет", "no", "false", "0", "nofaol", "нофаол"}
MAX_ROWS = 20000


class ImportFileError(Exception):
    """Fayl darajasidagi xato (format, xavfsizlik, majburiy ustun yo'q)."""


def norm_header(value) -> str:
    """Sarlavha: kichik harf, apostrof/bo'shliqlarsiz, kirill → lotin (solishtirish uchun)."""
    s = str(value or "").strip().lower()
    s = re.sub(r"[‘’ʻʼ'`´\s_\-]+", "", s)
    table = str.maketrans({
        "а": "a", "б": "b", "в": "v", "г": "g", "д": "d", "е": "e", "ё": "yo", "ж": "j", "з": "z", "и": "i",
        "й": "y", "к": "k", "л": "l", "м": "m", "н": "n", "о": "o", "п": "p", "р": "r", "с": "s", "т": "t",
        "у": "u", "ф": "f", "х": "x", "ц": "ts", "ч": "ch", "ш": "sh", "щ": "sh", "ъ": "", "ь": "", "э": "e",
        "ю": "yu", "я": "ya", "ў": "o", "қ": "q", "ғ": "g", "ҳ": "h",
    })
    return s.translate(table)


def parse_date(value) -> date | None:
    if value in (None, ""):
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    s = str(value).strip()
    for fmt in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            continue
    raise ValueError(f"Сана тушунарсиз: {s!r} (кк.оо.йййй ёки йййй-оо-кк)")


def parse_bool(value, default=True) -> bool:
    if value in (None, ""):
        return default
    s = str(value).strip().lower()
    if s in TRUE_WORDS:
        return True
    if s in FALSE_WORDS:
        return False
    raise ValueError(f"Ҳа/йўқ қиймати тушунарсиз: {value!r}")


def text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


@dataclass
class Sheet:
    headers: list
    rows: list  # (row_number, dict[normalized_header -> value])

    def col(self, row: dict, *aliases, required=False):
        for alias in aliases:
            key = norm_header(alias)
            if key in row and row[key] not in (None, ""):
                return row[key]
        if required:
            raise ValueError(f"«{aliases[0]}» устуни бўш")
        return None


def read_sheet(source, required_headers: dict) -> Sheet:
    """`source` — yo'l (str/Path) yoki fayl obyekti. Avval `excel_safety` tekshiruvi."""
    if isinstance(source, (str, Path)):
        path = Path(source)
        if not path.exists():
            raise ImportFileError(f"Файл топилмади: {source}")
        fh = path.open("rb")
    else:
        fh = source
    try:
        fh.seek(0)
        try:
            check_xlsx_zip(fh)
        except UnsafeExcelFile as exc:
            raise ImportFileError(f"Excel файл хавфсизлик текширувидан ўтмади: {exc}")
        fh.seek(0)
        try:
            wb = load_workbook(fh, read_only=True, data_only=True)
        except Exception as exc:  # noqa: BLE001 — buzilgan fayl
            raise ImportFileError(f"Excel файл ўқилмади: {exc}")
        ws = wb.active
        iterator = ws.iter_rows(values_only=True)
        try:
            header_row = next(iterator)
        except StopIteration:
            raise ImportFileError("Файл бўш")
        headers = [norm_header(h) for h in header_row]
        missing = [aliases[0] for aliases in required_headers.values() if not any(norm_header(a) in headers for a in aliases)]
        if missing:
            raise ImportFileError("Мажбурий устунлар топилмади: " + ", ".join(missing))
        rows = []
        for idx, values in enumerate(iterator, start=2):
            if values is None or all(v in (None, "") for v in values):
                continue
            rows.append((idx, {h: v for h, v in zip(headers, values) if h}))
            if len(rows) > MAX_ROWS:
                raise ImportFileError(f"Қаторлар жуда кўп (>{MAX_ROWS}) — файлни бўлинг")
        wb.close()
        return Sheet(headers=headers, rows=rows)
    finally:
        if isinstance(source, (str, Path)):
            fh.close()


@dataclass
class ImportResult:
    dry_run: bool
    partial: bool
    created: int = 0
    updated: int = 0
    unchanged: int = 0
    errors: list = field(default_factory=list)  # [{"row": n, "message": "..."}]
    applied: bool = False

    def error(self, row_no, message):
        self.errors.append({"row": row_no, "message": message})

    def to_dict(self) -> dict:
        return {"dry_run": self.dry_run, "applied": self.applied, "created": self.created, "updated": self.updated,
                "unchanged": self.unchanged, "errors": self.errors, "error_count": len(self.errors)}

    def lines(self) -> list[str]:
        prefix = "[DRY-RUN] " if self.dry_run else ""
        out = [f"{prefix}yangi: {self.created}, yangilandi: {self.updated}, o'zgarishsiz: {self.unchanged}, xato: {len(self.errors)}"]
        out += [f"  {e['row']}-qator: {e['message']}" for e in self.errors]
        return out


def message_of(exc) -> str:
    if isinstance(exc, ValidationError):
        if hasattr(exc, "message_dict"):
            return "; ".join(f"{k}: {', '.join(v)}" for k, v in exc.message_dict.items())
        return "; ".join(exc.messages)
    return str(getattr(exc, "message", None) or exc)


def _run(sheet: Sheet, result: ImportResult, handle_row):
    """Har qatorni savepoint ichida bajaradi; oxirida dry-run / xato bo'lsa hammasi bekor."""
    with transaction.atomic():
        for row_no, row in sheet.rows:
            try:
                with transaction.atomic():
                    handle_row(row_no, row)
            except (ValueError, ValidationError) as exc:
                result.error(row_no, message_of(exc))
            except Exception as exc:  # noqa: BLE001 — servis xatolari (TransitionError va h.k.)
                if hasattr(exc, "message"):
                    result.error(row_no, exc.message)
                else:
                    raise
        if result.dry_run or (result.errors and not result.partial):
            transaction.set_rollback(True)
            result.applied = False
        else:
            result.applied = True
    return result


def _save(obj, user):
    if user is not None:
        obj._history_user = user
    obj.save()


# ----------------------------------------------------------------------------- mahallalar (A1)
MAHALLA_HEADERS = {
    "source_code": ["манба коди", "manba kodi", "код", "kod"],
    "region_code": ["вилоят коди", "viloyat kodi"],
    "district_code": ["туман коди", "tuman kodi"],
    "name": ["маҳалла номи", "mahalla nomi", "номи", "nomi"],
}


def import_mahallas(source, *, dry_run=False, partial=True, user=None) -> ImportResult:
    from apps.billing.models import District, Region
    from apps.yashil.models import Mahalla

    H = MAHALLA_HEADERS
    sheet = read_sheet(source, H)
    result = ImportResult(dry_run=dry_run, partial=partial)
    regions = {r.code: r for r in Region.objects.all()}
    districts = {d.code: d for d in District.objects.select_related("region")}

    def handle(row_no, row):
        source_code = text(sheet.col(row, *H["source_code"], required=True))
        region_code = text(sheet.col(row, *H["region_code"], required=True))
        district_code = text(sheet.col(row, *H["district_code"], required=True))
        name = text(sheet.col(row, *H["name"], required=True))
        region, district = regions.get(region_code), districts.get(district_code)
        if region is None:
            raise ValueError(f"вилоят коди топилмади: {region_code}")
        if district is None or district.region_id != region.id:
            raise ValueError(f"туман коди топилмади ёки вилоятга мос эмас: {district_code}")
        values = {"name": name, "region": region, "district": district, "address": text(sheet.col(row, "манзил", "manzil")),
                  "is_active": parse_bool(sheet.col(row, "фаол", "faol"), default=True)}
        valid_from = parse_date(sheet.col(row, "амал бошланиши", "amal boshlanishi"))
        if valid_from:
            values["valid_from"] = valid_from
        values["valid_to"] = parse_date(sheet.col(row, "амал тугаши", "amal tugashi"))
        existing = Mahalla.objects.filter(source_code=source_code).first()
        if existing is None:
            _save(Mahalla(source_code=source_code, **values), user)
            result.created += 1
        elif any(getattr(existing, k) != v for k, v in values.items()):
            for k, v in values.items():
                setattr(existing, k, v)
            _save(existing, user)
            result.updated += 1
        else:
            result.unchanged += 1

    return _run(sheet, result, handle)


# ----------------------------------------------------------------------------- biriktirishlar (A2)
ASSIGNMENT_HEADERS = {
    "mahalla": ["маҳалла коди", "mahalla kodi", "маҳалла", "mahalla"],
    "organization": ["хўжалик", "xo'jalik", "xojalik", "ташкилот", "tashkilot"],
    "valid_from": ["амал бошланиши", "amal boshlanishi"],
}


def find_org(value: str):
    from apps.general.models import Organization

    try:
        return Organization.objects.get(pk=uuid.UUID(value))
    except (ValueError, Organization.DoesNotExist):
        pass
    org = Organization.objects.filter(name__iexact=value).first()
    if org is None:
        raise ValueError(f"ташкилот топилмади: {value}")
    return org


def _mahalla_by_code(code):
    from apps.yashil.models import Mahalla

    mahalla = Mahalla.objects.filter(source_code=code).first()
    if mahalla is None:
        raise ValueError(f"маҳалла коди топилмади: {code}")
    return mahalla


def import_assignments(source, *, dry_run=False, partial=True, user=None, source_doc="") -> ImportResult:
    from apps.yashil.models import SupplierAssignment

    H = ASSIGNMENT_HEADERS
    sheet = read_sheet(source, H)
    result = ImportResult(dry_run=dry_run, partial=partial)

    def handle(row_no, row):
        mahalla = _mahalla_by_code(text(sheet.col(row, *H["mahalla"], required=True)))
        org = find_org(text(sheet.col(row, *H["organization"], required=True)))
        valid_from = parse_date(sheet.col(row, *H["valid_from"], required=True))
        valid_to = parse_date(sheet.col(row, "амал тугаши", "amal tugashi"))
        existing = SupplierAssignment.objects.filter(mahalla=mahalla, organization=org, valid_from=valid_from).first()
        if existing is not None:
            if existing.valid_to != valid_to:
                existing.valid_to = valid_to
                _save(existing, user)
                result.updated += 1
            else:
                result.unchanged += 1
            return
        _save(SupplierAssignment(mahalla=mahalla, organization=org, valid_from=valid_from, valid_to=valid_to, source_doc=source_doc), user)
        result.created += 1

    return _run(sheet, result, handle)


# ----------------------------------------------------------------------------- mahalla → uyushma (A4)
ROUTE_HEADERS = {
    "mahalla": ["маҳалла коди", "mahalla kodi", "маҳалла", "mahalla"],
    "unit": ["уюшма бўлинмаси", "uyushma bo'linmasi", "уюшма", "uyushma"],
    "valid_from": ["амал бошланиши", "amal boshlanishi"],
}


def _find_unit(value: str):
    from apps.yashil.models import UnionUnit

    try:
        return UnionUnit.objects.get(pk=uuid.UUID(value))
    except (ValueError, UnionUnit.DoesNotExist):
        pass
    unit = UnionUnit.objects.filter(name__iexact=value).first()
    if unit is None:
        raise ValueError(f"уюшма бўлинмаси топилмади: {value}")
    return unit


def import_routes(source, *, dry_run=False, partial=True, user=None, source_doc="") -> ImportResult:
    from apps.yashil.models import MahallaUnionRoute

    H = ROUTE_HEADERS
    sheet = read_sheet(source, H)
    result = ImportResult(dry_run=dry_run, partial=partial)

    def handle(row_no, row):
        mahalla = _mahalla_by_code(text(sheet.col(row, *H["mahalla"], required=True)))
        unit = _find_unit(text(sheet.col(row, *H["unit"], required=True)))
        valid_from = parse_date(sheet.col(row, *H["valid_from"], required=True))
        valid_to = parse_date(sheet.col(row, "амал тугаши", "amal tugashi"))
        existing = MahallaUnionRoute.objects.filter(mahalla=mahalla, union_unit=unit, valid_from=valid_from).first()
        if existing is not None:
            if existing.valid_to != valid_to:
                existing.valid_to = valid_to
                _save(existing, user)
                result.updated += 1
            else:
                result.unchanged += 1
            return
        _save(MahallaUnionRoute(mahalla=mahalla, union_unit=unit, valid_from=valid_from, valid_to=valid_to, source_doc=source_doc), user)
        result.created += 1

    return _run(sheet, result, handle)


# ----------------------------------------------------------------------------- kalendar (Д1)
CALENDAR_HEADERS = {"date": ["сана", "sana", "date"], "kind": ["тури", "turi", "kind"]}
WORKDAY_WORDS = {norm_header(w) for w in ("иш куни", "ish kuni", "workday", "ишкуни", "кўчирилган иш куни")}
HOLIDAY_WORDS = {norm_header(w) for w in ("байрам", "bayram", "дам олиш", "dam olish", "holiday", "дам олиш куни")}


def import_calendar(source, *, dry_run=False, partial=True, user=None, source_ref="") -> ImportResult:
    from apps.yashil.models import CalendarDay

    H = CALENDAR_HEADERS
    sheet = read_sheet(source, H)
    result = ImportResult(dry_run=dry_run, partial=partial)

    def handle(row_no, row):
        day = parse_date(sheet.col(row, *H["date"], required=True))
        kind_raw = norm_header(text(sheet.col(row, *H["kind"], required=True)))
        if kind_raw in WORKDAY_WORDS:
            kind = CalendarDay.Kind.WORKDAY
        elif kind_raw in HOLIDAY_WORDS:
            kind = CalendarDay.Kind.HOLIDAY
        else:
            raise ValueError(f"тури тушунарсиз: {kind_raw!r} (байрам / иш куни)")
        values = {"kind": kind, "note": text(sheet.col(row, "изоҳ", "izoh", "note")), "source_ref": source_ref}
        existing = CalendarDay.objects.filter(date=day).first()
        if existing is None:
            _save(CalendarDay(date=day, **values), user)
            result.created += 1
        elif any(getattr(existing, k) != v for k, v in values.items()):
            for k, v in values.items():
                setattr(existing, k, v)
            _save(existing, user)
            result.updated += 1
        else:
            result.unchanged += 1

    return _run(sheet, result, handle)


# ----------------------------------------------------------------------------- boshlang'ich qoldiq (Е1)
OPENING_HEADERS = {
    "code": ["каталог коди", "katalog kodi", "код", "kod", "позиция коди", "pozitsiya kodi"],
    "qty": ["миқдор", "miqdor", "қолдиқ", "qoldiq"],
    "doc_number": ["ҳужжат рақами", "hujjat raqami", "асос ҳужжат", "asos hujjat"],
    "doc_date": ["ҳужжат санаси", "hujjat sanasi", "сана", "sana"],
}


def import_opening_stock(source, *, cabinet, dry_run=False, partial=False) -> ImportResult:
    """Yetkazib beruvchi omboriga boshlang'ich qoldiqlar (har pozitsiya bir marta, hujjat bilan)."""
    from apps.yashil.models import CatalogItem
    from apps.yashil.services import ledger

    H = OPENING_HEADERS
    sheet = read_sheet(source, H)
    result = ImportResult(dry_run=dry_run, partial=partial)
    warehouse = ledger.default_warehouse(cabinet.organization)
    seen = set()

    def handle(row_no, row):
        code = text(sheet.col(row, *H["code"], required=True))
        item = CatalogItem.objects.filter(code=code).first()
        if item is None:
            raise ValueError(f"каталог коди топилмади: {code}")
        if code in seen:
            raise ValueError(f"позиция файлда такрорланган: {code}")
        seen.add(code)
        raw = text(sheet.col(row, *H["qty"], required=True)).replace(" ", "").replace(",", ".")
        try:
            qty = Decimal(raw)
        except InvalidOperation:
            raise ValueError(f"миқдор нотўғри: {raw!r}")
        doc_number = text(sheet.col(row, *H["doc_number"], required=True))
        doc_date = parse_date(sheet.col(row, *H["doc_date"], required=True))
        ledger.post_opening(warehouse, item, qty, cabinet, doc_number=doc_number, doc_date=doc_date,
                            note=text(sheet.col(row, "изоҳ", "izoh", "note")))
        result.created += 1

    return _run(sheet, result, handle)
